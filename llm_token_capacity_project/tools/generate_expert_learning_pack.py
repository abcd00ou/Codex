"""Generate an expert Korean learning pack from high-quality references."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT / "outputs" / "reports"
RUN_DATE = "2026-05-15"


SOURCES = [
    ("IEA_AI", "Energy and AI", "IEA", "2025", "https://www.iea.org/reports/energy-and-ai", "Tier 2 institution", "global data center electricity demand, power supply, energy policy", "A01,A02,A03"),
    ("IEA_DEMAND", "Energy demand from AI", "IEA", "2025", "https://www.iea.org/reports/energy-and-ai/energy-demand-from-ai", "Tier 2 institution", "AI server power density and data center electricity demand", "A01,A02"),
    ("IEA_KEY_Q", "Key Questions on Energy and AI", "IEA", "2025", "https://www.iea.org/reports/key-questions-on-energy-and-ai/executive-summary", "Tier 2 institution", "AI power swings, storage, financial market interpretation", "A01,A02,A05"),
    ("LBNL_DC_2024", "2024 United States Data Center Energy Usage Report", "LBNL / DOE", "2024", "https://eta-publications.lbl.gov/publications/2024-lbnl-data-center-energy-usage-report", "Tier 2 government/lab", "US data center energy baseline and AI server modeling", "A01,A02,A03,A04"),
    ("CRS_DC_FAQ", "Data Centers and Their Energy Consumption: FAQ", "Congressional Research Service", "2026", "https://www.congress.gov/crs-product/R48646", "Tier 2 government", "policy framing and terminology", "A01,A02"),
    ("EPRI_POWERING", "Powering Intelligence", "EPRI", "2024", "https://www.epri.com/research/products/000000003002028905", "Tier 2 industry research", "US data center electricity scenarios and grid planning", "A01,A02"),
    ("EPRI_EPOCH_SCALING", "Scaling Intelligence", "EPRI / Epoch AI", "2025", "https://www.epri.com/research/products/000000003002033669", "Tier 2 research", "AI power capacity and frontier training power scenario", "A02,A06"),
    ("UPTIME_SURVEY_2024", "Global Data Center Survey 2024", "Uptime Institute", "2024", "https://uptimeinstitute.com/resources/research-and-reports/uptime-institute-global-data-center-survey-results-2024", "Tier 2 industry research", "PUE, cooling, power constraints, operator survey", "A03,A04"),
    ("UPTIME_COOLING_2024", "Cooling Systems Survey 2024", "Uptime Institute", "2024", "https://intelligence.uptimeinstitute.com/sites/default/files/2024-05/Uptime%20Institute%20Cooling%20Systems%20Survey%202024_0.pdf", "Tier 2 industry research", "liquid cooling and cooling architecture", "A03"),
    ("NREL_PUE", "High-Performance Computing Data Center PUE", "NREL", "accessed 2026-05-15", "https://www.nrel.gov/computational-science/measuring-efficiency-pue", "Tier 2 lab", "PUE concept and HPC efficiency", "A03"),
    ("GOOGLE_PUE", "Power usage effectiveness", "Google Data Centers", "accessed 2026-05-15", "https://datacenters.google/efficiency/", "Tier 1 company", "company PUE disclosure and efficiency framing", "A03"),
    ("META_DC", "Meta Sustainability: Data centers", "Meta", "accessed 2026-05-15", "https://sustainability.atmeta.com/data-centers/", "Tier 1 company", "data center sustainability and PUE context", "A03"),
    ("KAPLAN", "Scaling Laws for Neural Language Models", "Kaplan et al.", "2020", "https://arxiv.org/abs/2001.08361", "Tier 2 paper", "training compute scaling laws", "A06,A07"),
    ("CHINCHILLA", "Training Compute-Optimal Large Language Models", "Hoffmann et al.", "2022", "https://arxiv.org/abs/2203.15556", "Tier 2 paper", "compute-optimal model/data allocation", "A06,A07"),
    ("SEVILLA_COMPUTE", "Compute Trends Across Three Eras of Machine Learning", "Sevilla et al.", "2022", "https://arxiv.org/abs/2202.05924", "Tier 2 paper", "historical training compute trends", "A06"),
    ("EPOCH_TRAIN", "Training compute of frontier AI models grows by 4-5x per year", "Epoch AI", "2024", "https://epoch.ai/blog/training-compute-of-frontier-ai-models-grows-by-4-5x-per-year", "Tier 2 research", "frontier training compute growth", "A06"),
    ("EPOCH_POWER", "How much power will frontier AI training demand in 2030?", "Epoch AI", "2025", "https://epoch.ai/blog/power-demands-of-frontier-ai-training", "Tier 2 research", "training run power envelope", "A02,A06"),
    ("EPOCH_SCALING_2030", "Can AI scaling continue through 2030?", "Epoch AI", "2025", "https://epoch.ai/blog/can-ai-scaling-continue-through-2030", "Tier 2 research", "scaling bottlenecks, power/data/capex", "A06"),
    ("STANFORD_INDEX", "AI Index Report 2025", "Stanford HAI", "2025", "https://hai.stanford.edu/ai-index/2025-ai-index-report", "Tier 2 institution", "AI capability, cost, compute, policy", "A06,A08,A10"),
    ("DEEPSEEK_V3", "DeepSeek-V3 GitHub / Technical Report", "DeepSeek", "2024", "https://github.com/deepseek-ai/DeepSeek-V3", "Tier 1 model source", "MoE total and active parameter anchor", "A07,A08"),
    ("DEEPSEEK_R1", "DeepSeek-R1", "DeepSeek / arXiv", "2025", "https://arxiv.org/abs/2501.12948", "Tier 1/2 model source", "reasoning model and MoE parameter context", "A07,A08"),
    ("QWEN3", "Qwen3 GitHub", "Alibaba Qwen", "2025", "https://github.com/QwenLM/Qwen3", "Tier 1 model source", "MoE active parameter and model family", "A07,A08"),
    ("OPENAI_MODELS", "OpenAI model docs", "OpenAI", "accessed 2026-05-15", "https://platform.openai.com/docs/models", "Tier 1 company", "closed model disclosure boundary and product surface", "A07,A10"),
    ("PAGED_ATTENTION", "Efficient Memory Management for LLM Serving with PagedAttention", "Kwon et al.", "2023", "https://arxiv.org/abs/2309.06180", "Tier 2 paper", "KV cache, memory management, vLLM serving throughput", "A08,A09"),
    ("SPLITWISE", "Splitwise: Efficient generative LLM inference using phase splitting", "Patel et al.", "2023", "https://arxiv.org/abs/2311.18677", "Tier 2 paper", "prefill/decode phase splitting", "A08,A09"),
    ("MS_SPLITWISE", "Splitwise Microsoft Research blog", "Microsoft Research", "2024", "https://www.microsoft.com/en-us/research/blog/splitwise-improves-gpu-usage-by-splitting-llm-inference-phases/", "Tier 2 research blog", "operator-friendly phase split explanation", "A08,A09"),
    ("DISTSERVE", "DistServe: Disaggregating Prefill and Decoding", "Zhong et al.", "2024", "https://arxiv.org/abs/2401.09670", "Tier 2 paper", "TTFT/TPOT SLO, goodput, prefill/decode disaggregation", "A08,A09"),
    ("PD_DISAGG", "Prefill-Decode Aggregation or Disaggregation?", "arXiv", "2025", "https://arxiv.org/abs/2508.01989", "Tier 2 paper", "when prefill/decode disaggregation helps", "A08,A09"),
    ("ENERGY_INFERENCE", "Energy Use of AI Inference", "arXiv / Cell Patterns", "2025/2026", "https://arxiv.org/abs/2509.20241", "Tier 2 paper", "joules/token, test-time compute and inference energy", "A08,A09"),
    ("TOKEN_POWER_BENCH", "TokenPowerBench", "AAAI/arXiv", "2026", "https://arxiv.org/abs/2512.03024", "Tier 2 paper", "LLM inference power benchmarking", "A08,A09"),
    ("INFERENCEX", "InferenceX / InferenceMAX", "SemiAnalysis", "accessed 2026-05-15", "https://inferencex.semianalysis.com/", "Proxy/benchmark", "inference benchmark and tokens/MW sanity check", "A08,A09"),
    ("GOOGLE_IRONWOOD", "Ironwood: TPU for the age of inference", "Google Cloud", "2025", "https://blog.google/products/google-cloud/ironwood-tpu-age-of-inference/", "Tier 1 company", "inference-oriented TPU and perf/W direction", "A05,A08"),
    ("GOOGLE_TPU_V5E", "Cloud TPU v5e documentation", "Google Cloud", "accessed 2026-05-15", "https://docs.cloud.google.com/tpu/docs/v5e", "Tier 1 docs", "TPU training/serving deployment concepts", "A05,A08,A09"),
    ("NVIDIA_H100", "NVIDIA H100 Tensor Core GPU", "NVIDIA", "accessed 2026-05-15", "https://www.nvidia.com/en-us/data-center/h100/", "Tier 1 company", "Hopper baseline for accelerator assumptions", "A08"),
    ("NVIDIA_H200", "NVIDIA H200 Tensor Core GPU", "NVIDIA", "accessed 2026-05-15", "https://www.nvidia.com/en-us/data-center/h200/", "Tier 1 company", "HBM capacity/bandwidth uplift", "A08"),
    ("NVIDIA_GB200", "NVIDIA GB200 NVL72", "NVIDIA", "accessed 2026-05-15", "https://www.nvidia.com/en-us/data-center/gb200-nvl72/", "Tier 1 company", "rack-scale Blackwell/NVL72 architecture", "A02,A03,A08"),
    ("NVIDIA_DGX_GB", "NVIDIA DGX GB Rack Scale Systems User Guide", "NVIDIA Docs", "accessed 2026-05-15", "https://docs.nvidia.com/dgx/dgxgb200-user-guide/hardware.html", "Tier 1 docs", "rack-scale power, liquid cooling, system architecture", "A02,A03"),
    ("AWS_RAINIER", "AWS Project Rainier", "Amazon", "2025", "https://www.aboutamazon.com/news/aws/aws-project-rainier-ai-trainium-chips-compute-cluster", "Tier 1 company", "Anthropic/AWS Trainium2 cluster and attribution", "A02,A06,A10"),
    ("AWS_TRN2", "Amazon EC2 Trn2", "AWS", "accessed 2026-05-15", "https://aws.amazon.com/ec2/instance-types/trn2/", "Tier 1 docs", "Trainium2 instance context", "A08,A10"),
    ("AWS_INF2", "Amazon EC2 Inf2", "AWS", "accessed 2026-05-15", "https://aws.amazon.com/ec2/instance-types/inf2/", "Tier 1 docs", "Inferentia2 inference context", "A08"),
    ("OPENAI_INFRA", "Building compute infrastructure for the intelligence age", "OpenAI", "2025", "https://openai.com/index/building-the-compute-infrastructure-for-the-intelligence-age/", "Tier 1 company", "OpenAI compute infrastructure and capacity framing", "A01,A02,A10"),
    ("OPENAI_STARGATE", "Five new Stargate sites", "OpenAI", "2025", "https://openai.com/index/five-new-stargate-sites/", "Tier 1 company", "planned capacity and host/model-owner attribution", "A01,A02,A10"),
    ("SEMI_DC_MODEL", "Datacenter Industry Model", "SemiAnalysis", "accessed 2026-05-15", "https://newsletter.semianalysis.com/p/datacenter-model", "Market/technical research", "critical IT power and AI accelerator deployment modeling", "A01,A02,A04"),
    ("SEMI_DC_ANATOMY", "Datacenter Anatomy: Electrical Systems", "SemiAnalysis", "accessed 2026-05-15", "https://newsletter.semianalysis.com/p/datacenter-anatomy-part-1-electrical", "Market/technical research", "electrical chain and critical IT power", "A01,A02,A03"),
    ("SEMI_100K_H100", "100,000 H100 Clusters", "SemiAnalysis", "accessed 2026-05-15", "https://newsletter.semianalysis.com/p/100000-h100-clusters-power-network", "Market/technical research", "cluster power, networking, checkpointing", "A02,A04,A06"),
    ("EPOCH_GPU_POWER", "GPUs account for about 40% of power usage in AI data centers", "Epoch AI", "2025", "https://epoch.ai/data-insights/gpus-power-usage-in-ai-data-centers", "Tier 2 research", "facility/IT/GPU power decomposition", "A03,A04"),
    ("DELOITTE_AI_DC", "AI infrastructure gaps", "Deloitte Insights", "2025", "https://www.deloitte.com/us/en/insights/industry/power-and-utilities/data-center-infrastructure-artificial-intelligence.html", "Market context", "AI infrastructure, power/grid survey context", "A01,A02,A05"),
    ("MCK_POWER_COOL", "Beyond compute: infrastructure that powers and cools AI data centers", "McKinsey", "2025", "https://www.mckinsey.com/industries/industrials-and-electronics/our-insights/beyond-compute-infrastructure-that-powers-and-cools-ai-data-centers", "Market context", "power/cooling investment and constraints", "A02,A03"),
    ("MCK_WORKLOADS", "The next big shifts in AI workloads and hyperscaler strategies", "McKinsey", "2025", "https://www.mckinsey.com/industries/technology-media-and-telecommunications/our-insights/the-next-big-shifts-in-ai-workloads-and-hyperscaler-strategies", "Market context", "workload shift and hyperscaler strategy", "A05,A06,A10"),
    ("BAIN_AI_COMPUTE", "AI Changes Big and Small Computing", "Bain", "2024", "https://www.bain.com/insights/ai-changes-big-and-small-computing-tech-report-2024/", "Market context", "AI computing architecture and supply-chain framing", "A04,A05,A10"),
    ("PWC_SEMI", "Semiconductor and Beyond 2026", "PwC", "2024", "https://www.pwc.com/gx/en/industries/technology/pwc-semiconductor-and-beyond-2026-full-report.pdf", "Market context", "semiconductor, HBM and datacenter server context", "A07,A08"),
    ("AI_2027", "AI 2027", "AI Futures Project", "2025", "https://ai-2027.com/", "Scenario/stress test", "aggressive capability and compute-demand shock scenario", "A05,A06,A08,A09,A10"),
]


MODULES = [
    (
        "1. 데이터센터 전력과 grid를 먼저 이해한다",
        "계약 GW, active GW, facility power, IT load를 구분하지 못하면 이후 token forecast는 전부 흔들립니다. 이 모듈은 IEA, LBNL, EPRI, Uptime, NREL을 중심으로 전력 수요와 데이터센터 물리 구조를 공부합니다.",
        ["IEA_AI", "IEA_DEMAND", "LBNL_DC_2024", "EPRI_POWERING", "UPTIME_SURVEY_2024", "NREL_PUE"],
        ["contracted_power_gw와 active_power_gw를 분리해서 설명할 수 있다.", "PUE가 무엇을 의미하고 무엇을 의미하지 않는지 설명할 수 있다.", "facility power, IT load, accelerator board power를 구분할 수 있다."],
    ),
    (
        "2. training compute는 scaling law에서 출발한다",
        "frontier training reserve를 이해하려면 Kaplan, Chinchilla, Epoch AI를 읽어야 합니다. training은 단기 token output이 아니라 미래 capability와 모델 세대 전환을 위한 투자입니다.",
        ["KAPLAN", "CHINCHILLA", "SEVILLA_COMPUTE", "EPOCH_TRAIN", "EPOCH_POWER", "EPOCH_SCALING_2030"],
        ["training compute와 model/data scaling 관계를 설명할 수 있다.", "training run peak power와 연중 평균 training share를 구분할 수 있다.", "training_power_share를 단순 잔여값 이상으로 해석할 수 있다."],
    ),
    (
        "3. inference serving은 prefill/decode와 KV cache가 핵심이다",
        "상용 LLM economics는 inference serving에서 결정됩니다. vLLM, Splitwise, DistServe는 tokens/sec/MW와 utilization을 이해하는 핵심 논문입니다.",
        ["PAGED_ATTENTION", "SPLITWISE", "MS_SPLITWISE", "DISTSERVE", "PD_DISAGG", "ENERGY_INFERENCE", "TOKEN_POWER_BENCH"],
        ["prefill과 decode의 차이를 설명할 수 있다.", "TTFT, TPOT, goodput, KV cache가 utilization에 주는 영향을 설명할 수 있다.", "benchmark와 production 평균 throughput을 구분할 수 있다."],
    ),
    (
        "4. hardware와 rack-scale system은 token capacity의 물리적 바닥이다",
        "GPU/TPU/Trainium/Inferentia는 FLOPs뿐 아니라 HBM, interconnect, rack power, cooling, software stack이 다릅니다. official hardware docs를 benchmark보다 먼저 읽어야 합니다.",
        ["GOOGLE_IRONWOOD", "GOOGLE_TPU_V5E", "NVIDIA_H100", "NVIDIA_H200", "NVIDIA_GB200", "NVIDIA_DGX_GB", "AWS_RAINIER", "AWS_TRN2", "AWS_INF2"],
        ["accelerator generation과 rack-scale system의 차이를 설명할 수 있다.", "board TDP와 facility power를 구분할 수 있다.", "custom ASIC/TPU가 tokens/MW 가정에 주는 영향을 설명할 수 있다."],
    ),
    (
        "5. company/model-owner attribution을 따로 배운다",
        "OpenAI, Microsoft, Anthropic, AWS, Oracle, Google Cloud가 얽히면 같은 capacity가 여러 번 계산됩니다. attribution rule은 forecast의 회계 원칙입니다.",
        ["OPENAI_INFRA", "OPENAI_STARGATE", "AWS_RAINIER", "OPENAI_MODELS", "MCK_WORKLOADS", "DELOITTE_AI_DC"],
        ["model owner, host provider, product owner를 구분할 수 있다.", "OpenAI/Microsoft와 Anthropic/AWS 중복 계산 위험을 설명할 수 있다.", "attribution_rule을 보고서 문구에 반영할 수 있다."],
    ),
    (
        "6. market/consulting/source는 scenario context로 읽는다",
        "Deloitte, McKinsey, Bain, PwC, SemiAnalysis, AI 2027은 매우 유용하지만 source class가 다릅니다. 숫자 fact가 아니라 context, scenario, proxy, stress test로 분류해야 합니다.",
        ["SEMI_DC_MODEL", "SEMI_DC_ANATOMY", "SEMI_100K_H100", "EPOCH_GPU_POWER", "DELOITTE_AI_DC", "MCK_POWER_COOL", "MCK_WORKLOADS", "BAIN_AI_COMPUTE", "PWC_SEMI", "AI_2027", "INFERENCEX"],
        ["시장 리포트를 fact로 오용하지 않는 법을 설명할 수 있다.", "benchmark/proxy/context/scenario를 분리할 수 있다.", "AI-2027 stress scenario를 Base와 분리해 다룰 수 있다."],
    ),
]


def source_map() -> dict[str, tuple[str, str, str, str, str, str, str, str]]:
    return {row[0]: row for row in SOURCES}


def source_table(rows: list[tuple[str, str, str, str, str, str, str, str]]) -> list[str]:
    lines = ["| ID | 자료 | 발행자 | 연도 | Source class | 핵심 사용처 | Agents |", "|---|---|---|---|---|---|---|"]
    for sid, title, publisher, year, url, klass, use, agents in rows:
        lines.append(f"| {sid} | [{title}]({url}) | {publisher} | {year} | {klass} | {use} | {agents} |")
    return lines


def build_markdown() -> str:
    smap = source_map()
    lines = [
        "# Expert Learning Pack: LLM Compute, Power, Inference, and Assumptions",
        "",
        f"생성일: {RUN_DATE}",
        "",
        "## 목적",
        "",
        "이 자료는 LLM token capacity simulation을 더 전문적으로 운영하기 위한 학습 패키지입니다. 단순히 출처를 많이 모으는 것이 아니라, 각 자료가 어떤 가정을 개선하는지, 어떤 자료는 fact이고 어떤 자료는 proxy/context/scenario인지 구분하는 능력을 키우는 데 목적이 있습니다.",
        "",
        "## 학습의 핵심 원칙",
        "",
        "1. 전력 숫자는 반드시 `contracted -> active -> IT load -> AI IT load -> inference/training` 순서로 내려옵니다.",
        "2. training compute와 inference compute는 같은 accelerator를 쓸 수 있어도 운영 목적과 scheduling 방식이 다릅니다.",
        "3. inference는 prefill, decode, KV cache, latency SLO, batching, utilization의 문제입니다.",
        "4. hardware peak FLOPs는 production tokens/sec/MW가 아닙니다.",
        "5. consulting report와 scenario report는 매우 유용하지만 company-level production fact로 직접 쓰지 않습니다.",
        "6. 모든 숫자는 source class와 derivation type을 달고 agent evidence로 승격해야 합니다.",
        "",
        "## 6-Module Study Plan",
        "",
    ]
    for idx, (title, intro, source_ids, outcomes) in enumerate(MODULES, 1):
        lines.extend([f"### Module {idx}. {title}", "", intro, "", "**핵심 읽기 자료**", ""])
        for sid in source_ids:
            _, stitle, publisher, year, url, klass, use, agents = smap[sid]
            lines.append(f"- **{sid}:** [{stitle}]({url}) - {publisher}, {year}. `{klass}`. 사용처: {use}. Agents: {agents}.")
        lines.extend(["", "**학습 후 할 수 있어야 하는 것**", ""])
        for outcome in outcomes:
            lines.append(f"- {outcome}")
        lines.extend(
            [
                "",
                "**실습 과제**",
                "",
                "1. 위 자료 중 하나를 골라 원문 숫자 3개를 추출합니다.",
                "2. 각 숫자를 Fact, Derived Estimate, Proxy, Scenario 중 하나로 분류합니다.",
                "3. 해당 assumption agent의 `evidence.md`에 넣을 수 있는 evidence row 초안을 작성합니다.",
                "4. 이 source가 Base를 바꾸는지, Bear/Bull만 바꾸는지, 별도 stress scenario인지 판단합니다.",
                "",
            ]
        )

    lines.extend(
        [
            "## Expert Reading Notes",
            "",
            "### Power and Data Center Reports",
            "",
            "IEA, LBNL, EPRI, Uptime 계열 자료는 회사별 token forecast를 직접 주지 않습니다. 대신 전력 수요, grid 제약, data center energy accounting, PUE, cooling, AI server deployment의 물리적 경계 조건을 제공합니다. 이 자료를 읽을 때는 항상 scope를 확인해야 합니다. 글로벌 데이터센터 전력인지, 미국 데이터센터 전력인지, AI-specific power인지, 전체 IT load인지가 다르면 같은 숫자처럼 비교할 수 없습니다.",
            "",
            "LBNL과 IEA 자료는 `A01 contracted_power_gw`, `A02 active_power_gw`, `A03 pue`, `A04 ai_workload_share`의 이론적 바닥을 만듭니다. 하지만 특정 OpenAI/Google/Meta active GW를 직접 제공하지 않으므로, agent evidence에 넣을 때는 macro boundary 또는 mechanism evidence로 분류해야 합니다.",
            "",
            "### Training Compute Literature",
            "",
            "Kaplan과 Chinchilla는 training compute를 이해하는 언어를 제공합니다. Kaplan은 model/data/compute scaling의 직관을 제공했고, Chinchilla는 고정 compute에서 model size와 training tokens의 균형을 재정의했습니다. Epoch AI와 EPRI/Epoch는 이 이론이 frontier training power demand로 어떻게 이어지는지 보여줍니다. 단, training run의 peak power envelope를 연중 평균 training_power_share로 직접 바꾸면 안 됩니다.",
            "",
            "### Inference Serving Literature",
            "",
            "PagedAttention, Splitwise, DistServe는 LLM inference가 단순히 GPU FLOPs 문제가 아니라 memory management, prefill/decode 분리, latency SLO, goodput, scheduling의 문제임을 보여줍니다. 이 세 논문은 `A08 tokens_per_second_per_mw`와 `A09 utilization`의 필수 배경입니다. 특히 benchmark throughput과 production average throughput을 분리해서 읽어야 합니다.",
            "",
            "### Energy-to-Token Research",
            "",
            "Energy Use of AI Inference, TokenPowerBench, recent joules/token benchmark들은 우리가 tokens/sec/MW만 보지 말고 joules/token, active binding constraint, utilization-adjusted output을 같이 봐야 한다는 방향을 줍니다. 이 자료들은 아직 회사별 production telemetry가 아니므로 proxy/benchmark로 분류합니다.",
            "",
            "### Hardware and Rack-Scale Documents",
            "",
            "NVIDIA, Google, AWS official docs는 hardware generation, HBM, interconnect, rack-scale system, Trainium/Inferentia/TPU 같은 custom accelerator의 1차 자료입니다. 이 자료에서 얻는 것은 peak capability와 architecture context입니다. production tokens/MW는 별도 serving efficiency와 utilization을 거쳐야 합니다.",
            "",
            "### Market and Scenario Reports",
            "",
            "Deloitte, McKinsey, Bain, PwC, SemiAnalysis, AI 2027은 방향성과 scenario를 읽는 데 매우 유용합니다. 그러나 source class가 다릅니다. SemiAnalysis는 technical market research와 benchmark/proxy로, Deloitte/McKinsey/Bain/PwC는 market context로, AI 2027은 aggressive stress scenario로 분류해야 합니다.",
            "",
            "## Full Annotated Source Library",
            "",
            *source_table(SOURCES),
            "",
            "## Agent Promotion Checklist",
            "",
            "source를 실제 모델에 쓰기 전에 다음 절차를 따릅니다.",
            "",
            "1. `data/source_review_log.md`에 source를 기록합니다.",
            "2. 관련 agent의 `evidence.md`에 source-reviewed evidence row를 추가합니다.",
            "3. evidence_class를 Fact, Derived Estimate, Proxy, Scenario로 표시합니다.",
            "4. 숫자 단위가 GW, MW, TWh/year, tokens/sec, tokens/day, parameters 중 무엇인지 명시합니다.",
            "5. `state.md`에 변경 후보를 적고, active power/inference share/tokens/MW/utilization/attribution이면 orchestrator review를 받습니다.",
            "",
            "## Suggested 4-Week Study Sprint",
            "",
            "### Week 1: Power and Active Capacity",
            "",
            "- Read IEA, LBNL, EPRI, Uptime.",
            "- Output: A01/A02/A03 evidence candidates.",
            "",
            "### Week 2: Training and Scaling",
            "",
            "- Read Kaplan, Chinchilla, Epoch AI, EPRI/Epoch.",
            "- Output: A06 training reserve scenario notes.",
            "",
            "### Week 3: Inference Serving and Tokens/MW",
            "",
            "- Read PagedAttention, Splitwise, DistServe, Energy Use of AI Inference, InferenceX.",
            "- Output: A08/A09 benchmark/proxy framework.",
            "",
            "### Week 4: Hardware, Market Context, Attribution",
            "",
            "- Read Google Ironwood, NVIDIA GB200/DGX GB, AWS Rainier, OpenAI infrastructure, Deloitte/McKinsey/Bain/PwC, AI 2027.",
            "- Output: A05/A10 scenario and attribution notes.",
        ]
    )
    return "\n".join(lines)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    run.font.name = "Apple SD Gothic Neo"
    run.font.size = Pt(10)


def build_docx(markdown_text: str, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    styles["Normal"].font.name = "Apple SD Gothic Neo"
    styles["Normal"].font.size = Pt(10)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Apple SD Gothic Neo"
        styles[style_name].font.color.rgb = RGBColor(15, 23, 42)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Expert Learning Pack")
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = "Apple SD Gothic Neo"
    run.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"LLM Compute, Power, Inference, and Assumptions | {RUN_DATE}")

    for line in markdown_text.splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("|"):
            # Keep tables readable in DOCX as plain text to avoid very wide tables.
            add_para(doc, line)
        else:
            add_para(doc, line.replace("**", "").replace("`", ""))

    doc.save(path)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    markdown = build_markdown()
    md_path = DOCS / "expert_learning_pack.md"
    docx_path = OUT / "expert_learning_pack_kr.docx"
    md_path.write_text(markdown, encoding="utf-8")
    build_docx(markdown, docx_path)
    print({"status": "PASS", "sources": len(SOURCES), "markdown": str(md_path), "docx": str(docx_path)})


if __name__ == "__main__":
    main()
