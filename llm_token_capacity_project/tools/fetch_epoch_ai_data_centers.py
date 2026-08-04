"""Fetch and summarize Epoch AI data center data by AI user and supplier.

Outputs:
- raw ZIP and extracted CSVs
- normalized user-site rows
- user-year, supplier-year, and supplier-user-year power summaries for 2026-2030
- Markdown and XLSX summary reports
"""

from __future__ import annotations

import csv
import re
import shutil
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data" / "epoch_ai_data_centers"
RAW_DIR = DATA_DIR / "raw"
NORM_DIR = DATA_DIR / "normalized"
OUT_DIR = DATA_DIR / "outputs"
ZIP_URL = "https://epoch.ai/data/data_centers/data_centers.zip"
RUN_DATE = date.today().isoformat()
YEARS = list(range(2026, 2031))


@dataclass
class UserTag:
    name: str
    confidence: str


def ensure_dirs() -> None:
    for path in (RAW_DIR, NORM_DIR, OUT_DIR):
        path.mkdir(parents=True, exist_ok=True)


def download_zip() -> Path:
    zip_path = RAW_DIR / f"epoch_ai_data_centers_{RUN_DATE}.zip"
    latest_path = RAW_DIR / "epoch_ai_data_centers_latest.zip"
    if not zip_path.exists():
        if latest_path.exists():
            zip_path = latest_path
        else:
            archived = sorted(RAW_DIR.glob("epoch_ai_data_centers_*.zip"), key=lambda path: path.stat().st_mtime, reverse=True)
            if archived:
                zip_path = archived[0]
            else:
                with urllib.request.urlopen(ZIP_URL, timeout=60) as response:
                    zip_path.write_bytes(response.read())
    if zip_path != latest_path:
        shutil.copyfile(zip_path, latest_path)
    return zip_path


def extract_zip(zip_path: Path) -> dict[str, Path]:
    extracted: dict[str, Path] = {}
    with ZipFile(zip_path) as archive:
        for name in archive.namelist():
            if name.endswith("/"):
                continue
            target = RAW_DIR / name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(archive.read(name))
            extracted[Path(name).name] = target
    return extracted


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, rows: list[dict[str, object]], headers: list[str]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def parse_float(value: str | None) -> float:
    if value is None:
        return 0.0
    cleaned = value.replace(",", "").strip()
    if not cleaned:
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def parse_date(value: str) -> date | None:
    if not value:
        return None
    try:
        return datetime.strptime(value[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def parse_tagged_people(value: str) -> list[UserTag]:
    if not value.strip():
        return []
    people: list[UserTag] = []
    for part in re.split(r",\s*", value):
        item = part.strip()
        if not item:
            continue
        match = re.match(r"^(.*?)\s+#(confident|likely|speculative)\s*$", item)
        if match:
            people.append(UserTag(match.group(1).strip(), match.group(2)))
        else:
            people.append(UserTag(item, "unlabeled"))
    return people


def latest_at_or_before(rows: list[dict[str, str]], cutoff: date) -> dict[str, str] | None:
    candidates = [row for row in rows if (row_date := parse_date(row.get("Date", ""))) and row_date <= cutoff]
    if not candidates:
        return None
    return max(candidates, key=lambda row: parse_date(row.get("Date", "")) or date.min)


def first_reaches_capacity(rows: list[dict[str, str]], field: str, target: float) -> date | None:
    if target <= 0:
        return None
    dated_rows = sorted(
        [(parse_date(row.get("Date", "")), row) for row in rows],
        key=lambda pair: pair[0] or date.max,
    )
    for row_date, row in dated_rows:
        if row_date and parse_float(row.get(field)) >= target * 0.999:
            return row_date
    return None


def confidence_rank(confidence: str) -> int:
    return {"confident": 3, "likely": 2, "speculative": 1, "unlabeled": 0}.get(confidence, 0)


def blank_year_summary(entity_key: str, entity_name: str, year: int, confidence_key: str, confidence: str) -> dict[str, object]:
    return {
        entity_key: entity_name,
        "year": year,
        "site_count": 0,
        confidence_key: confidence,
        "it_power_mw_full_exposure": 0.0,
        "it_power_mw_equal_split": 0.0,
        "facility_power_mw_full_exposure": 0.0,
        "facility_power_mw_equal_split": 0.0,
        "h100_equiv_full_exposure": 0.0,
        "h100_equiv_equal_split": 0.0,
        "performance_8bit_ops_full_exposure": 0.0,
        "performance_8bit_ops_equal_split": 0.0,
        "total_cost_2025_usd_bn_full_exposure": 0.0,
        "total_cost_2025_usd_bn_equal_split": 0.0,
        "current_it_power_mw_full_exposure": 0.0,
        "current_it_power_mw_equal_split": 0.0,
        "current_h100_equiv_full_exposure": 0.0,
        "current_h100_equiv_equal_split": 0.0,
        "current_total_cost_2025_usd_bn_full_exposure": 0.0,
        "current_total_cost_2025_usd_bn_equal_split": 0.0,
        "contracted_or_planned_it_power_mw_full_exposure": 0.0,
        "contracted_or_planned_it_power_mw_equal_split": 0.0,
        "contracted_or_planned_facility_power_mw_full_exposure": 0.0,
        "contracted_or_planned_facility_power_mw_equal_split": 0.0,
        "contracted_or_planned_h100_equiv_full_exposure": 0.0,
        "contracted_or_planned_h100_equiv_equal_split": 0.0,
        "contracted_or_planned_8bit_ops_full_exposure": 0.0,
        "contracted_or_planned_8bit_ops_equal_split": 0.0,
        "earliest_first_operational_date": "",
        "latest_completion_date_for_max_it_power": "",
        "latest_timeline_date_used": "",
        "data_centers": [],
        "notes": [],
    }


def year_values_for_site(site_timelines: list[dict[str, str]], year: int) -> tuple[dict[str, float], str, str]:
    cutoff = date(year, 12, 31)
    year_row = latest_at_or_before(site_timelines, cutoff)
    if year_row is None:
        return (
            {
                "it_power_mw": 0.0,
                "facility_power_mw": 0.0,
                "h100_equivalents": 0.0,
                "performance_8bit_ops": 0.0,
                "total_cost_2025_usd_bn": 0.0,
            },
            "",
            "",
        )
    return (
        {
            "it_power_mw": parse_float(year_row.get("IT power (MW)")),
            "facility_power_mw": parse_float(year_row.get("Power (MW)")),
            "h100_equivalents": parse_float(year_row.get("H100 equivalents")),
            "performance_8bit_ops": parse_float(year_row.get("Performance (8-bit OP/s)")),
            "total_cost_2025_usd_bn": parse_float(year_row.get("Total capital cost (2025 USD billions)")),
        },
        year_row.get("Date", ""),
        year_row.get("Construction status", ""),
    )


def apply_site_year_to_summary(
    summary: dict[str, object],
    confidence_key: str,
    confidence: str,
    values: dict[str, float],
    split_count: int,
    max_it_power: float,
    max_power: float,
    max_h100: float,
    max_perf: float,
    current_it_power: float,
    current_h100: float,
    current_cost: float,
    first_operational: date | None,
    completion: date | None,
    timeline_date: str,
    status: str,
    site_name: str,
) -> None:
    summary["site_count"] = int(summary["site_count"]) + 1
    if confidence_rank(confidence) > confidence_rank(str(summary[confidence_key])):
        summary[confidence_key] = confidence
    for metric, value in values.items():
        exposure_key = {
            "it_power_mw": "it_power_mw_full_exposure",
            "facility_power_mw": "facility_power_mw_full_exposure",
            "h100_equivalents": "h100_equiv_full_exposure",
            "performance_8bit_ops": "performance_8bit_ops_full_exposure",
            "total_cost_2025_usd_bn": "total_cost_2025_usd_bn_full_exposure",
        }[metric]
        split_key = exposure_key.replace("_full_exposure", "_equal_split")
        summary[exposure_key] = float(summary[exposure_key]) + value
        summary[split_key] = float(summary[split_key]) + value / split_count
    summary["contracted_or_planned_it_power_mw_full_exposure"] = float(summary["contracted_or_planned_it_power_mw_full_exposure"]) + max_it_power
    summary["contracted_or_planned_it_power_mw_equal_split"] = float(summary["contracted_or_planned_it_power_mw_equal_split"]) + max_it_power / split_count
    summary["contracted_or_planned_facility_power_mw_full_exposure"] = float(summary["contracted_or_planned_facility_power_mw_full_exposure"]) + max_power
    summary["contracted_or_planned_facility_power_mw_equal_split"] = float(summary["contracted_or_planned_facility_power_mw_equal_split"]) + max_power / split_count
    summary["contracted_or_planned_h100_equiv_full_exposure"] = float(summary["contracted_or_planned_h100_equiv_full_exposure"]) + max_h100
    summary["contracted_or_planned_h100_equiv_equal_split"] = float(summary["contracted_or_planned_h100_equiv_equal_split"]) + max_h100 / split_count
    summary["contracted_or_planned_8bit_ops_full_exposure"] = float(summary["contracted_or_planned_8bit_ops_full_exposure"]) + max_perf
    summary["contracted_or_planned_8bit_ops_equal_split"] = float(summary["contracted_or_planned_8bit_ops_equal_split"]) + max_perf / split_count
    summary["current_it_power_mw_full_exposure"] = float(summary["current_it_power_mw_full_exposure"]) + current_it_power
    summary["current_it_power_mw_equal_split"] = float(summary["current_it_power_mw_equal_split"]) + current_it_power / split_count
    summary["current_h100_equiv_full_exposure"] = float(summary["current_h100_equiv_full_exposure"]) + current_h100
    summary["current_h100_equiv_equal_split"] = float(summary["current_h100_equiv_equal_split"]) + current_h100 / split_count
    summary["current_total_cost_2025_usd_bn_full_exposure"] = float(summary["current_total_cost_2025_usd_bn_full_exposure"]) + current_cost
    summary["current_total_cost_2025_usd_bn_equal_split"] = float(summary["current_total_cost_2025_usd_bn_equal_split"]) + current_cost / split_count
    if first_operational:
        current = str(summary["earliest_first_operational_date"])
        if not current or first_operational.isoformat() < current:
            summary["earliest_first_operational_date"] = first_operational.isoformat()
    if completion:
        current = str(summary["latest_completion_date_for_max_it_power"])
        if not current or completion.isoformat() > current:
            summary["latest_completion_date_for_max_it_power"] = completion.isoformat()
    if timeline_date and timeline_date > str(summary["latest_timeline_date_used"]):
        summary["latest_timeline_date_used"] = timeline_date
    summary["data_centers"].append(site_name)
    if status:
        summary["notes"].append(f"{site_name} @ {timeline_date}: {status}")


def finalize_year_rows(rows: list[dict[str, object]], sort_keys: tuple[str, ...]) -> list[dict[str, object]]:
    final_rows: list[dict[str, object]] = []
    for row in rows:
        out = row.copy()
        out["data_centers"] = "; ".join(sorted(set(out["data_centers"])))
        out["notes"] = " | ".join(out["notes"][:5])
        for key, value in list(out.items()):
            if isinstance(value, float):
                out[key] = round(value, 6)
        final_rows.append(out)
    final_rows.sort(key=lambda row: tuple(row[key] if key != "year" else int(row[key]) for key in sort_keys))
    return final_rows


def summarize(data_centers: list[dict[str, str]], timelines: list[dict[str, str]]) -> tuple[
    list[dict[str, object]],
    list[dict[str, object]],
    list[dict[str, object]],
    list[dict[str, object]],
    list[dict[str, object]],
    list[dict[str, object]],
]:
    timeline_by_site: dict[str, list[dict[str, str]]] = {}
    for row in timelines:
        timeline_by_site.setdefault(row["Data center"], []).append(row)

    user_site_rows: list[dict[str, object]] = []
    supplier_site_rows: list[dict[str, object]] = []
    user_year: dict[tuple[str, int], dict[str, object]] = {}
    supplier_year: dict[tuple[str, int], dict[str, object]] = {}
    supplier_user_year: dict[tuple[str, str, int], dict[str, object]] = {}
    site_summary: list[dict[str, object]] = []

    for site in data_centers:
        site_name = site["Name"]
        site_timelines = sorted(
            timeline_by_site.get(site_name, []),
            key=lambda row: parse_date(row.get("Date", "")) or date.min,
        )
        users = parse_tagged_people(site.get("Users", ""))
        owners = parse_tagged_people(site.get("Owner", ""))
        user_count = max(len(users), 1)
        max_it_power = max([parse_float(row.get("IT power (MW)")) for row in site_timelines] + [parse_float(site.get("Current power (MW)"))])
        max_power = max([parse_float(row.get("Power (MW)")) for row in site_timelines] + [0.0])
        max_h100 = max([parse_float(row.get("H100 equivalents")) for row in site_timelines] + [parse_float(site.get("Current H100 equivalents"))])
        max_perf = max([parse_float(row.get("Performance (8-bit OP/s)")) for row in site_timelines] + [0.0])
        max_cost = max([parse_float(row.get("Total capital cost (2025 USD billions)")) for row in site_timelines] + [parse_float(site.get("Current total capital cost (2025 USD billions)"))])
        current_it_power = parse_float(site.get("Current power (MW)"))
        current_h100 = parse_float(site.get("Current H100 equivalents"))
        current_cost = parse_float(site.get("Current total capital cost (2025 USD billions)"))
        completion = first_reaches_capacity(site_timelines, "IT power (MW)", max_it_power)
        first_operational = first_reaches_capacity(site_timelines, "IT power (MW)", 1.0)
        latest_timeline_date = max([parse_date(row.get("Date", "")) for row in site_timelines if parse_date(row.get("Date", ""))] or [None])
        owner_names = "; ".join(f"{owner.name} #{owner.confidence}" for owner in owners)

        site_summary.append(
            {
                "data_center": site_name,
                "owner": owner_names,
                "users": site.get("Users", ""),
                "country": site.get("Country", ""),
                "project": site.get("Project", ""),
                "current_it_power_mw": current_it_power,
                "contracted_or_planned_it_power_mw": round(max_it_power, 3),
                "contracted_or_planned_facility_power_mw": round(max_power, 3),
                "contracted_or_planned_h100_equivalents": round(max_h100, 3),
                "contracted_or_planned_8bit_ops": round(max_perf, 3),
                "contracted_or_planned_total_cost_2025_usd_bn": round(max_cost, 6),
                "first_operational_date": first_operational.isoformat() if first_operational else "",
                "completion_date_for_max_it_power": completion.isoformat() if completion else "",
                "latest_timeline_date": latest_timeline_date.isoformat() if latest_timeline_date else "",
                "current_chip_types": site.get("Current chip types", ""),
                "all_chip_types": site.get("All chip types", ""),
                "selected_sources": site.get("Selected Sources", ""),
                "calculations_sheet": site.get("Calculations sheet", ""),
            }
        )

        if not users:
            users = [UserTag("Unknown user", "unlabeled")]
        user_site_owner_names = owner_names
        if not owners:
            owners = [UserTag("Unknown supplier", "unlabeled")]
        owner_count = max(len(owners), 1)
        for user in users:
            user_site_rows.append(
                {
                    "user": user.name,
                    "user_confidence": user.confidence,
                    "data_center": site_name,
                    "owner": user_site_owner_names,
                    "country": site.get("Country", ""),
                    "project": site.get("Project", ""),
                    "current_it_power_mw_full_exposure": current_it_power,
                    "current_it_power_mw_equal_split": current_it_power / user_count,
                    "current_h100_equiv_full_exposure": current_h100,
                    "current_h100_equiv_equal_split": current_h100 / user_count,
                    "current_total_cost_2025_usd_bn_full_exposure": current_cost,
                    "current_total_cost_2025_usd_bn_equal_split": current_cost / user_count,
                    "contracted_or_planned_it_power_mw_full_exposure": round(max_it_power, 3),
                    "contracted_or_planned_it_power_mw_equal_split": round(max_it_power / user_count, 3),
                    "contracted_or_planned_facility_power_mw_full_exposure": round(max_power, 3),
                    "contracted_or_planned_facility_power_mw_equal_split": round(max_power / user_count, 3),
                    "contracted_or_planned_h100_equiv_full_exposure": round(max_h100, 3),
                    "contracted_or_planned_h100_equiv_equal_split": round(max_h100 / user_count, 3),
                    "contracted_or_planned_8bit_ops_full_exposure": round(max_perf, 3),
                    "contracted_or_planned_8bit_ops_equal_split": round(max_perf / user_count, 3),
                    "completion_date_for_max_it_power": completion.isoformat() if completion else "",
                    "first_operational_date": first_operational.isoformat() if first_operational else "",
                    "latest_timeline_date": latest_timeline_date.isoformat() if latest_timeline_date else "",
                    "current_chip_types": site.get("Current chip types", ""),
                    "all_chip_types": site.get("All chip types", ""),
                    "selected_sources": site.get("Selected Sources", ""),
                    "calculations_sheet": site.get("Calculations sheet", ""),
                }
            )

            for year in YEARS:
                values, timeline_date, status = year_values_for_site(site_timelines, year)
                key = (user.name, year)
                summary = user_year.setdefault(
                    key,
                    blank_year_summary("user", user.name, year, "highest_user_confidence", user.confidence),
                )
                apply_site_year_to_summary(
                    summary,
                    "highest_user_confidence",
                    user.confidence,
                    values,
                    user_count,
                    max_it_power,
                    max_power,
                    max_h100,
                    max_perf,
                    current_it_power,
                    current_h100,
                    current_cost,
                    first_operational,
                    completion,
                    timeline_date,
                    status,
                    site_name,
                )

        for owner in owners:
            supplier_site_rows.append(
                {
                    "supplier": owner.name,
                    "supplier_confidence": owner.confidence,
                    "data_center": site_name,
                    "users": site.get("Users", ""),
                    "country": site.get("Country", ""),
                    "project": site.get("Project", ""),
                    "current_it_power_mw_full_exposure": current_it_power,
                    "current_it_power_mw_equal_split": current_it_power / owner_count,
                    "current_h100_equiv_full_exposure": current_h100,
                    "current_h100_equiv_equal_split": current_h100 / owner_count,
                    "current_total_cost_2025_usd_bn_full_exposure": current_cost,
                    "current_total_cost_2025_usd_bn_equal_split": current_cost / owner_count,
                    "contracted_or_planned_it_power_mw_full_exposure": round(max_it_power, 3),
                    "contracted_or_planned_it_power_mw_equal_split": round(max_it_power / owner_count, 3),
                    "contracted_or_planned_facility_power_mw_full_exposure": round(max_power, 3),
                    "contracted_or_planned_facility_power_mw_equal_split": round(max_power / owner_count, 3),
                    "contracted_or_planned_h100_equiv_full_exposure": round(max_h100, 3),
                    "contracted_or_planned_h100_equiv_equal_split": round(max_h100 / owner_count, 3),
                    "contracted_or_planned_8bit_ops_full_exposure": round(max_perf, 3),
                    "contracted_or_planned_8bit_ops_equal_split": round(max_perf / owner_count, 3),
                    "completion_date_for_max_it_power": completion.isoformat() if completion else "",
                    "first_operational_date": first_operational.isoformat() if first_operational else "",
                    "latest_timeline_date": latest_timeline_date.isoformat() if latest_timeline_date else "",
                    "current_chip_types": site.get("Current chip types", ""),
                    "all_chip_types": site.get("All chip types", ""),
                    "selected_sources": site.get("Selected Sources", ""),
                    "calculations_sheet": site.get("Calculations sheet", ""),
                }
            )
            for year in YEARS:
                values, timeline_date, status = year_values_for_site(site_timelines, year)
                supplier_summary = supplier_year.setdefault(
                    (owner.name, year),
                    blank_year_summary("supplier", owner.name, year, "highest_supplier_confidence", owner.confidence),
                )
                apply_site_year_to_summary(
                    supplier_summary,
                    "highest_supplier_confidence",
                    owner.confidence,
                    values,
                    owner_count,
                    max_it_power,
                    max_power,
                    max_h100,
                    max_perf,
                    current_it_power,
                    current_h100,
                    current_cost,
                    first_operational,
                    completion,
                    timeline_date,
                    status,
                    site_name,
                )
                for user in users:
                    bridge_summary = supplier_user_year.setdefault(
                        (owner.name, user.name, year),
                        {
                            **blank_year_summary("supplier", owner.name, year, "highest_supplier_confidence", owner.confidence),
                            "user": user.name,
                            "highest_user_confidence": user.confidence,
                        },
                    )
                    if confidence_rank(user.confidence) > confidence_rank(str(bridge_summary["highest_user_confidence"])):
                        bridge_summary["highest_user_confidence"] = user.confidence
                    apply_site_year_to_summary(
                        bridge_summary,
                        "highest_supplier_confidence",
                        owner.confidence,
                        values,
                        owner_count * user_count,
                        max_it_power,
                        max_power,
                        max_h100,
                        max_perf,
                        current_it_power,
                        current_h100,
                        current_cost,
                        first_operational,
                        completion,
                        timeline_date,
                        status,
                        site_name,
                    )

    year_rows = finalize_year_rows(list(user_year.values()), ("user", "year"))
    supplier_year_rows = finalize_year_rows(list(supplier_year.values()), ("supplier", "year"))
    supplier_user_year_rows = finalize_year_rows(list(supplier_user_year.values()), ("supplier", "user", "year"))
    bridge_front = ["supplier", "user", "year", "site_count", "highest_supplier_confidence", "highest_user_confidence"]
    supplier_user_year_rows = [
        {**{key: row[key] for key in bridge_front}, **{key: value for key, value in row.items() if key not in bridge_front}}
        for row in supplier_user_year_rows
    ]
    user_site_rows.sort(key=lambda row: (str(row["user"]), str(row["data_center"])))
    supplier_site_rows.sort(key=lambda row: (str(row["supplier"]), str(row["data_center"])))
    site_summary.sort(key=lambda row: str(row["data_center"]))
    return site_summary, user_site_rows, year_rows, supplier_site_rows, supplier_year_rows, supplier_user_year_rows


def write_markdown(
    site_rows: list[dict[str, object]],
    user_site_rows: list[dict[str, object]],
    user_year_rows: list[dict[str, object]],
    supplier_site_rows: list[dict[str, object]],
    supplier_year_rows: list[dict[str, object]],
    supplier_user_year_rows: list[dict[str, object]],
) -> Path:
    path = OUT_DIR / "epoch_ai_data_centers_user_summary_2026_2030.md"
    rows_2030 = [row for row in user_year_rows if row["year"] == 2030]
    supplier_rows_2030 = [row for row in supplier_year_rows if row["year"] == 2030]
    top_power = sorted(rows_2030, key=lambda row: float(row["it_power_mw_equal_split"]), reverse=True)[:15]
    top_contract = sorted(rows_2030, key=lambda row: float(row["contracted_or_planned_it_power_mw_equal_split"]), reverse=True)[:15]
    top_suppliers = sorted(supplier_rows_2030, key=lambda row: float(row["it_power_mw_equal_split"]), reverse=True)[:15]
    lines = [
        "# Epoch AI Data Centers User And Supplier Power Summary, 2026-2030",
        "",
        f"- Generated: {RUN_DATE}",
        f"- Source: {ZIP_URL}",
        "- Basis: Epoch AI AI Data Centers ZIP. Site-level users and owners are expanded into user-site, supplier-site, and supplier-user-year bridge rows.",
        "- Counting modes: `full_exposure` assigns every listed user/supplier the full site capacity; `equal_split` divides site capacity across listed users, suppliers, or supplier-user pairs to avoid double-counting in totals.",
        "- Contracted/planned capacity: maximum site timeline capacity observed through the downloaded dataset, not necessarily a legally contracted power-purchase amount.",
        "- Current capacity: Epoch site-level `Current power (MW)` and `Current H100 equivalents`, aggregated by listed user.",
        "",
        "## 2030 Top Users By Equal-Split IT Power",
        "",
        "| Rank | User | Sites | Current IT MW | 2030 IT Power MW | Contracted/Planned IT MW | 2030 H100-eq | Latest completion date |",
        "|---:|---|---:|---:|---:|---:|---:|---|",
    ]
    for idx, row in enumerate(top_power, 1):
        lines.append(
            f"| {idx} | {row['user']} | {row['site_count']} | {float(row['current_it_power_mw_equal_split']):,.0f} | "
            f"{float(row['it_power_mw_equal_split']):,.0f} | {float(row['contracted_or_planned_it_power_mw_equal_split']):,.0f} | "
            f"{float(row['h100_equiv_equal_split']):,.0f} | {row['latest_completion_date_for_max_it_power']} |"
        )
    lines += [
        "",
        "## 2030 Top Users By Contracted/Planned Equal-Split IT Power",
        "",
        "| Rank | User | Sites | Current IT MW | Contracted/Planned IT MW | 2030 IT Power MW | Contracted/Planned H100-eq | Data centers |",
        "|---:|---|---:|---:|---:|---:|---:|---|",
    ]
    for idx, row in enumerate(top_contract, 1):
        sites = str(row["data_centers"])
        if len(sites) > 120:
            sites = sites[:117] + "..."
        lines.append(
            f"| {idx} | {row['user']} | {row['site_count']} | {float(row['current_it_power_mw_equal_split']):,.0f} | "
            f"{float(row['contracted_or_planned_it_power_mw_equal_split']):,.0f} | "
            f"{float(row['it_power_mw_equal_split']):,.0f} | {float(row['contracted_or_planned_h100_equiv_equal_split']):,.0f} | {sites} |"
        )
    lines += [
        "",
        "## 2030 Top Suppliers By Equal-Split IT Power",
        "",
        "| Rank | Supplier | Sites | Current IT MW | 2030 IT Power MW | Contracted/Planned IT MW | 2030 H100-eq | Latest completion date |",
        "|---:|---|---:|---:|---:|---:|---:|---|",
    ]
    for idx, row in enumerate(top_suppliers, 1):
        lines.append(
            f"| {idx} | {row['supplier']} | {row['site_count']} | {float(row['current_it_power_mw_equal_split']):,.0f} | "
            f"{float(row['it_power_mw_equal_split']):,.0f} | {float(row['contracted_or_planned_it_power_mw_equal_split']):,.0f} | "
            f"{float(row['h100_equiv_equal_split']):,.0f} | {row['latest_completion_date_for_max_it_power']} |"
        )
    lines += [
        "",
        "## Output Files",
        "",
        "- `normalized/epoch_ai_data_center_sites.csv`: one row per data center.",
        "- `normalized/epoch_ai_data_center_user_sites.csv`: one row per user-data-center relationship.",
        "- `normalized/epoch_ai_data_center_user_year_summary_2026_2030.csv`: user-year summary table.",
        "- `normalized/epoch_ai_data_center_supplier_sites.csv`: one row per supplier-data-center relationship.",
        "- `normalized/epoch_ai_data_center_supplier_year_summary_2026_2030.csv`: supplier-year summary table.",
        "- `normalized/epoch_ai_data_center_supplier_user_year_bridge_2026_2030.csv`: supplier-user-year bridge table.",
        "- `outputs/epoch_ai_data_centers_user_year_power_2026_2030.csv`: dashboard-ready user-year power table.",
        "- `outputs/epoch_ai_data_centers_supplier_year_power_2026_2030.csv`: dashboard-ready supplier-year power table.",
        "- `outputs/epoch_ai_data_centers_supplier_user_year_power_2026_2030.csv`: dashboard-ready supplier-user-year power bridge.",
        "- `outputs/epoch_ai_data_centers_user_summary_2026_2030.xlsx`: workbook with the same tables.",
        "",
        "## Method Notes",
        "",
        "- Epoch states that the dataset covers construction timelines from satellite imagery, permits and public documents, and includes timeline estimates for IT power, compute and cost.",
        "- Epoch says 80% of IT power estimates are expected within a factor of 1.4x; compute within 1.5x; timeline estimates within 6 months.",
        "- User tags from Epoch are preserved as `#confident`, `#likely`, `#speculative`, or `unlabeled`.",
        "- Multi-user sites can be analyzed with full exposure for relationship mapping or equal split for non-double-counted aggregates.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_xlsx(
    site_rows: list[dict[str, object]],
    user_site_rows: list[dict[str, object]],
    user_year_rows: list[dict[str, object]],
    supplier_site_rows: list[dict[str, object]],
    supplier_year_rows: list[dict[str, object]],
    supplier_user_year_rows: list[dict[str, object]],
) -> Path:
    path = OUT_DIR / "epoch_ai_data_centers_user_summary_2026_2030.xlsx"
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return path

    wb = Workbook()
    wb.remove(wb.active)
    tables = [
        ("user_year_2026_2030", user_year_rows),
        ("supplier_year_2026_2030", supplier_year_rows),
        ("supplier_user_bridge", supplier_user_year_rows),
        ("user_site_map", user_site_rows),
        ("supplier_site_map", supplier_site_rows),
        ("site_summary", site_rows),
    ]
    for sheet_name, rows in tables:
        ws = wb.create_sheet(sheet_name)
        if not rows:
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="14213D")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx, header in enumerate(headers, 1):
            col_letter = get_column_letter(col_idx)
            width = min(max(len(header) + 2, 14), 56)
            if header in {"notes", "selected_sources", "data_centers"}:
                width = 72
            ws.column_dimensions[col_letter].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return path


def main() -> None:
    ensure_dirs()
    zip_path = download_zip()
    extracted = extract_zip(zip_path)
    data_centers = read_csv(extracted["data_centers.csv"])
    timelines = read_csv(extracted["data_center_timelines.csv"])
    site_rows, user_site_rows, user_year_rows, supplier_site_rows, supplier_year_rows, supplier_user_year_rows = summarize(data_centers, timelines)

    site_headers = list(site_rows[0].keys()) if site_rows else []
    user_site_headers = list(user_site_rows[0].keys()) if user_site_rows else []
    user_year_headers = list(user_year_rows[0].keys()) if user_year_rows else []
    supplier_site_headers = list(supplier_site_rows[0].keys()) if supplier_site_rows else []
    supplier_year_headers = list(supplier_year_rows[0].keys()) if supplier_year_rows else []
    supplier_user_year_headers = list(supplier_user_year_rows[0].keys()) if supplier_user_year_rows else []

    write_csv(NORM_DIR / "epoch_ai_data_center_sites.csv", site_rows, site_headers)
    write_csv(NORM_DIR / "epoch_ai_data_center_user_sites.csv", user_site_rows, user_site_headers)
    write_csv(NORM_DIR / "epoch_ai_data_center_user_year_summary_2026_2030.csv", user_year_rows, user_year_headers)
    write_csv(NORM_DIR / "epoch_ai_data_center_supplier_sites.csv", supplier_site_rows, supplier_site_headers)
    write_csv(NORM_DIR / "epoch_ai_data_center_supplier_year_summary_2026_2030.csv", supplier_year_rows, supplier_year_headers)
    write_csv(NORM_DIR / "epoch_ai_data_center_supplier_user_year_bridge_2026_2030.csv", supplier_user_year_rows, supplier_user_year_headers)
    output_user_year = OUT_DIR / "epoch_ai_data_centers_user_year_power_2026_2030.csv"
    output_supplier_year = OUT_DIR / "epoch_ai_data_centers_supplier_year_power_2026_2030.csv"
    output_supplier_user_year = OUT_DIR / "epoch_ai_data_centers_supplier_user_year_power_2026_2030.csv"
    write_csv(output_user_year, user_year_rows, user_year_headers)
    write_csv(output_supplier_year, supplier_year_rows, supplier_year_headers)
    write_csv(output_supplier_user_year, supplier_user_year_rows, supplier_user_year_headers)
    md_path = write_markdown(
        site_rows,
        user_site_rows,
        user_year_rows,
        supplier_site_rows,
        supplier_year_rows,
        supplier_user_year_rows,
    )
    xlsx_path = write_xlsx(
        site_rows,
        user_site_rows,
        user_year_rows,
        supplier_site_rows,
        supplier_year_rows,
        supplier_user_year_rows,
    )

    print(
        {
            "status": "PASS",
            "zip": str(zip_path),
            "sites": len(site_rows),
            "user_site_rows": len(user_site_rows),
            "user_year_rows": len(user_year_rows),
            "supplier_site_rows": len(supplier_site_rows),
            "supplier_year_rows": len(supplier_year_rows),
            "supplier_user_year_rows": len(supplier_user_year_rows),
            "outputs": [
                str(NORM_DIR / "epoch_ai_data_center_sites.csv"),
                str(NORM_DIR / "epoch_ai_data_center_user_sites.csv"),
                str(NORM_DIR / "epoch_ai_data_center_user_year_summary_2026_2030.csv"),
                str(NORM_DIR / "epoch_ai_data_center_supplier_sites.csv"),
                str(NORM_DIR / "epoch_ai_data_center_supplier_year_summary_2026_2030.csv"),
                str(NORM_DIR / "epoch_ai_data_center_supplier_user_year_bridge_2026_2030.csv"),
                str(output_user_year),
                str(output_supplier_year),
                str(output_supplier_user_year),
                str(md_path),
                str(xlsx_path),
            ],
        }
    )


if __name__ == "__main__":
    main()
