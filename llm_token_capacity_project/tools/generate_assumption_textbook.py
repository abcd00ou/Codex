"""Generate long-form Korean assumption textbooks for LLM token simulation.

Each assumption gets its own Markdown and Word report. The reports are written
as study material for building realistic, evidence-aware forecast assumptions.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "assumptions"
OUT = ROOT / "outputs" / "reports" / "assumptions"
RUN_DATE = "2026-05-15"


SOURCES = {
    "IEA_AI": ("IEA, Energy and AI", "IEA", "2025", "https://www.iea.org/reports/energy-and-ai"),
    "LBNL_DC": ("2024 United States Data Center Energy Usage Report", "LBNL / DOE", "2024", "https://buildings.lbl.gov/publications/2024-lbnl-data-center-energy-usage-report"),
    "CRS_DC": ("Data Centers and Their Energy Consumption: FAQ", "Congressional Research Service", "2026", "https://www.congress.gov/crs-product/R48646"),
    "UPTIME_2024": ("Global Data Center Survey 2024", "Uptime Institute", "2024", "https://uptimeinstitute.com/resources/research-and-reports/uptime-institute-global-data-center-survey-results-2024"),
    "GOOGLE_PUE": ("Power usage effectiveness", "Google Data Centers", "accessed 2026-05-15", "https://datacenters.google/efficiency/"),
    "NREL_PUE": ("High-Performance Computing Data Center PUE", "NREL", "accessed 2026-05-15", "https://www.nrel.gov/computational-science/measuring-efficiency-pue"),
    "META_SUST": ("Meta Sustainability: Data centers", "Meta", "accessed 2026-05-15", "https://sustainability.atmeta.com/data-centers/"),
    "OPENAI_INFRA": ("Building the compute infrastructure for the intelligence age", "OpenAI", "2025", "https://openai.com/index/building-the-compute-infrastructure-for-the-intelligence-age/"),
    "OPENAI_STARGATE": ("Five new Stargate sites and planned capacity", "OpenAI", "2025", "https://openai.com/index/five-new-stargate-sites/"),
    "GOOGLE_IRONWOOD": ("Ironwood: the first Google TPU for the age of inference", "Google Cloud", "2025", "https://blog.google/products/google-cloud/ironwood-tpu-age-of-inference/"),
    "GOOGLE_TPU_V5E": ("Cloud TPU v5e documentation", "Google Cloud", "accessed 2026-05-15", "https://docs.cloud.google.com/tpu/docs/v5e"),
    "AWS_RAINIER": ("AWS activates Project Rainier", "Amazon", "2025", "https://www.aboutamazon.com/news/aws/aws-project-rainier-ai-trainium-chips-compute-cluster"),
    "AWS_INF2": ("Amazon EC2 Inf2 instances", "AWS", "accessed 2026-05-15", "https://aws.amazon.com/ec2/instance-types/inf2/"),
    "AWS_TRN2": ("Amazon EC2 Trn2 instances", "AWS", "accessed 2026-05-15", "https://aws.amazon.com/ec2/instance-types/trn2/"),
    "NVIDIA_H100": ("NVIDIA H100 Tensor Core GPU", "NVIDIA", "accessed 2026-05-15", "https://www.nvidia.com/en-us/data-center/h100/"),
    "NVIDIA_H200": ("NVIDIA H200 Tensor Core GPU", "NVIDIA", "accessed 2026-05-15", "https://www.nvidia.com/en-us/data-center/h200/"),
    "NVIDIA_GB200": ("NVIDIA GB200 NVL72", "NVIDIA", "accessed 2026-05-15", "https://www.nvidia.com/en-us/data-center/gb200-nvl72/"),
    "KAPLAN": ("Scaling Laws for Neural Language Models", "Kaplan et al., arXiv", "2020", "https://arxiv.org/abs/2001.08361"),
    "CHINCHILLA": ("Training Compute-Optimal Large Language Models", "Hoffmann et al., arXiv", "2022", "https://arxiv.org/abs/2203.15556"),
    "EPOCH_TRAIN": ("Training compute of frontier AI models grows by 4-5x per year", "Epoch AI", "2024", "https://epoch.ai/blog/training-compute-of-frontier-ai-models-grows-by-4-5x-per-year"),
    "EPOCH_POWER": ("How much power will frontier AI training demand in 2030?", "Epoch AI", "2025", "https://epoch.ai/blog/power-demands-of-frontier-ai-training/"),
    "OPENAI_GPT41": ("GPT-4.1 model documentation", "OpenAI", "accessed 2026-05-15", "https://platform.openai.com/docs/models/gpt-4.1"),
    "DEEPSEEK_V3": ("DeepSeek-V3 GitHub / Technical Report", "DeepSeek", "2024", "https://github.com/deepseek-ai/DeepSeek-V3"),
    "DEEPSEEK_R1": ("DeepSeek-R1: Incentivizing Reasoning Capability", "DeepSeek, arXiv", "2025", "https://arxiv.org/abs/2501.12948"),
    "QWEN3": ("Qwen3 GitHub", "Alibaba Qwen", "2025", "https://github.com/QwenLM/Qwen3"),
    "PAGED_ATTENTION": ("Efficient Memory Management for LLM Serving with PagedAttention", "Kwon et al., arXiv", "2023", "https://arxiv.org/abs/2309.06180"),
    "SPLITWISE": ("Splitwise: Efficient generative LLM inference using phase splitting", "Patel et al., arXiv", "2023", "https://arxiv.org/abs/2311.18677"),
    "DISTSERVE": ("DistServe: Disaggregating Prefill and Decoding", "Zhong et al., arXiv", "2024", "https://arxiv.org/abs/2401.09670"),
    "MS_SPLITWISE": ("Splitwise improves GPU usage by splitting LLM inference phases", "Microsoft Research", "2024", "https://www.microsoft.com/en-us/research/blog/splitwise-improves-gpu-usage-by-splitting-llm-inference-phases/"),
}


COMMON_WARNING = (
    "이 리포트의 숫자 범위는 학습용 starting band입니다. 회사별 공식 telemetry가 공개되면 반드시 교체해야 하며, "
    "공식 source가 없는 값은 fact가 아니라 estimate/proxy/scenario로 표시해야 합니다."
)


ASSUMPTIONS = [
    {
        "id": "A01",
        "slug": "contracted_power_gw",
        "title": "contracted_power_gw",
        "subtitle": "계약·발표 전력 capacity를 어떻게 읽을 것인가",
        "sources": ["IEA_AI", "LBNL_DC", "CRS_DC", "OPENAI_INFRA", "OPENAI_STARGATE"],
        "starting_band": "회사별 publicly announced/contracted/planned GW를 상한선으로 사용. 2026-2030 forecast에서는 active_power_gw와 분리.",
        "body": [
            ("정의", [
                "contracted_power_gw는 특정 model owner 또는 그 host partner가 확보했거나 확보한다고 발표한 전력 capacity의 상한값이다. 여기에는 전력구매계약, utility interconnection, colocation lease, 데이터센터 캠퍼스 planned capacity, cloud partner capacity commitment가 섞일 수 있다. 중요한 점은 이 값이 token capacity가 아니라는 것이다. 전력 capacity는 토큰을 만드는 물리적 필요조건일 뿐, accelerator가 설치되고 cluster가 service-ready 상태가 되어야 비로소 active compute가 된다.",
                "AI 인프라 뉴스에서 GW 숫자는 매우 강한 인상을 주지만, 그 GW가 무엇을 의미하는지 정의하지 않으면 모델을 망친다. 어떤 보도는 site total power를 말하고, 어떤 발표는 multi-year planned capacity를 말하며, 어떤 자료는 one training run의 peak power envelope를 말한다. contracted_power_gw는 이런 숫자 중 회사가 미래에 접근할 수 있는 전력 상한을 잡는 변수이며, active_power_gw, ai_it_load_gw, inference_gw와 같은 하위 변수와 절대 섞으면 안 된다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "전력 계약은 데이터센터 프로젝트의 시작점에 가깝다. utility가 power를 제공할 수 있어야 하고, 송전망과 변전소가 준비되어야 하며, 현장 배전, UPS, switchgear, backup generation, cooling plant, water 또는 liquid cooling loop가 모두 맞아야 한다. AI 데이터센터는 일반 enterprise data center보다 rack density가 높기 때문에 동일한 MW라도 설계 난이도가 다르다. GB200 NVL72 같은 rack-scale system은 rack 자체가 전력·냉각·네트워크 설계 단위로 바뀌기 때문에, utility 계약만으로 operational readiness를 판단할 수 없다.",
                "계약 전력은 또한 business option value를 갖는다. model owner는 미래 model generation, inference growth, enterprise SLA, training run 경쟁에 대비해 capacity를 선점한다. 따라서 발표 capacity가 모두 바로 사용될 필요는 없다. 기업은 전력과 부지를 먼저 확보해 병목을 줄이고, accelerator supply와 제품 수요가 따라오는 속도에 맞춰 active capacity를 올릴 수 있다. 시뮬레이션에서 contracted_power_gw가 중요한 이유는 미래 상한을 정하기 때문이고, active_power_gw가 중요한 이유는 실제 token generation을 제한하기 때문이다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "contracted_power_gw는 직접 발표가 있는 경우 그 숫자를 그대로 fact anchor로 둔다. 단, 문구가 planned인지 committed인지 operational인지 반드시 기록한다. 발표가 없으면 capex, data center lease, cloud agreement, utility interconnection queue, partner announcement를 조합해 estimate로만 둔다. company별 row에서는 single precise number보다 low/base/high band가 더 안전하다.",
                "계약 전력의 현실적 범위는 기업 성격에 따라 다르다. OpenAI처럼 Stargate와 Oracle capacity 발표가 있는 경우 model owner attribution을 명시해야 한다. Anthropic은 AWS/Google host capacity와 Claude token owner attribution을 분리해야 한다. Google과 Meta처럼 자체 데이터센터와 accelerator strategy가 큰 기업은 company capex와 data center sustainability disclosure가 유용하다. 중국 업체는 공개성이 낮으므로 confidence를 낮추되, 단순히 낮은 공개성을 낮은 capacity로 해석하면 안 된다.",
            ]),
            ("모델링 영향", [
                "contracted_power_gw는 forecast의 천장이다. active_power_gw가 이 값을 넘으면 모델 오류다. 하지만 token forecast에 직접 곱하면 안 된다. 올바른 순서는 contracted_power_gw에서 active_power_gw를 만들고, PUE와 AI workload share를 적용한 뒤, inference share와 tokens/sec/MW를 곱하는 것이다.",
                "이 값이 10% 증가해도 단기 token forecast는 거의 변하지 않을 수 있다. 왜냐하면 2026년 forecast는 active deployment가 병목일 수 있기 때문이다. 반대로 2030년 forecast에서는 contracted capacity가 부족하면 active_power_gw의 상한이 되므로 장기 token capacity를 강하게 제한한다.",
            ]),
            ("Hallucination 위험", [
                "가장 큰 위험은 planned GW를 active inference GW로 표현하는 것이다. 두 번째 위험은 host provider capacity를 model owner capacity로 중복 계산하는 것이다. 세 번째 위험은 annual electricity consumption TWh와 instantaneous power GW를 섞는 것이다. 보고서 문구에는 반드시 planned/contracted/operational을 구분해야 한다.",
            ]),
        ],
    },
    {
        "id": "A02",
        "slug": "active_power_gw",
        "title": "active_power_gw",
        "subtitle": "실제로 AI cluster가 쓸 수 있는 운영 전력을 추정하는 법",
        "sources": ["LBNL_DC", "CRS_DC", "IEA_AI", "GOOGLE_TPU_V5E", "AWS_RAINIER", "NVIDIA_GB200"],
        "starting_band": "2026 active/contracted 15-45%, 2030 active/contracted 45-80%를 학습용 band로 시작. site-level disclosure가 있으면 교체.",
        "body": [
            ("정의", [
                "active_power_gw는 발표 또는 계약된 전력 중 실제로 데이터센터 facility에 들어와 운영 가능한 상태가 된 전력이다. 이 값은 전력망 연결, 변전, 냉각, 랙 설치, accelerator delivery, cluster qualification을 통과한 capacity를 의미한다. token forecast에서는 contracted_power_gw보다 active_power_gw가 훨씬 직접적인 제한 조건이다.",
                "active_power_gw는 여전히 facility-level 전력일 수 있다. 따라서 이 값 자체가 GPU board power 또는 inference power는 아니다. PUE를 적용해 IT load로 바꾸고, AI workload share를 적용해 AI IT load를 만든 뒤, inference/training split을 적용해야 한다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "데이터센터의 active capacity는 건설 완료와 동일하지 않다. building shell이 완성되어도 utility energization, switchgear test, cooling commissioning, network turn-up, server burn-in, orchestration integration이 남는다. AI cluster는 일반 서버보다 failure domain이 크다. 수만 GPU training cluster는 일부 rack만 준비되어도 전체 cluster efficiency가 크게 떨어질 수 있다. inference fleet은 더 작은 단위로 배치할 수 있지만, latency SLO와 regional redundancy 때문에 역시 완전한 active capacity가 필요하다.",
                "Google Cloud TPU 문서는 training과 serving workload의 provisioning 차이를 보여준다. 같은 accelerator product라도 training job은 throughput과 availability에, serving job은 latency에 최적화된다. 이는 active_power_gw가 단순 전력 가동 여부를 넘어, 해당 cluster가 어떤 workload에 적합하게 provisioned 되었는지를 함께 봐야 함을 의미한다.",
                "AWS Project Rainier처럼 특정 partner를 위한 대형 cluster가 online이라는 발표는 active compute에 가까운 신호다. 그러나 발표가 'nearly half a million Trainium2 chips' 같은 accelerator count를 제공하더라도, 그 전체가 특정 시점에 inference로 쓰인다는 뜻은 아니다. active compute와 workload allocation은 분리해야 한다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "active_power_gw를 추정할 때는 milestone 기반 접근이 가장 안전하다. 첫째, announced/contracted capacity가 있는지 확인한다. 둘째, construction 또는 energized milestone이 있는지 본다. 셋째, accelerator delivery 또는 cluster online 발표가 있는지 본다. 넷째, 서비스 배포 또는 API availability 증가와 연결되는지 본다. 이 네 단계가 모두 확인되면 active 비율을 높일 수 있다.",
                "공개자료가 부족하면 deployment ratio를 시나리오로 둔다. 2026년에는 많은 AI capacity가 ramp 중일 수 있으므로 active/contracted ratio를 낮게 잡고, 2030년에는 프로젝트 완공과 accelerator delivery가 누적되며 비율이 올라가는 구조가 현실적이다. 단, 전력망 병목, transformer shortage, cooling constraint, permitting delay가 있으면 2030년에도 active conversion이 낮을 수 있다.",
            ]),
            ("모델링 영향", [
                "active_power_gw는 token forecast의 가장 큰 sensitivity driver 중 하나다. 모든 하위 계산이 active power에서 시작하기 때문이다. active_power_gw가 10% 오르면, PUE, AI workload share, inference share, tokens/sec/MW, utilization이 같다는 조건에서 token forecast도 거의 10% 오른다.",
                "하지만 active_power_gw는 confidence가 낮은 경우가 많다. 따라서 단일 base 값만 두면 보고서가 취약해진다. 반드시 bear/base/bull deployment ratio를 두고, site-level source가 나오면 교체할 replacement path를 명시해야 한다.",
            ]),
            ("Hallucination 위험", [
                "active_power_gw를 발표 GW와 같게 두는 것이 가장 위험하다. 두 번째 위험은 accelerator count에서 board power만 곱해 facility active power를 추정하는 것이다. board power에는 CPU, memory, networking, power conversion, cooling overhead가 빠져 있다. 세 번째 위험은 online cluster 발표를 연중 평균 active power로 그대로 쓰는 것이다.",
            ]),
        ],
    },
    {
        "id": "A03",
        "slug": "pue",
        "title": "pue",
        "subtitle": "facility power를 IT load로 바꾸는 핵심 계수",
        "sources": ["UPTIME_2024", "GOOGLE_PUE", "NREL_PUE", "META_SUST", "LBNL_DC"],
        "starting_band": "AI hyperscale 신규 facility 학습 band: 1.08-1.25. 보수적 전체 fleet 또는 미공개 site는 1.15-1.35.",
        "body": [
            ("정의", [
                "PUE는 Power Usage Effectiveness의 약자로, 데이터센터 전체 facility energy를 IT equipment energy로 나눈 값이다. PUE가 1.2라면 facility가 1.2MW를 소비할 때 IT equipment가 약 1MW를 소비한다는 뜻이다. AI token forecast에서는 active facility power를 IT load로 바꾸기 위해 사용한다.",
                "PUE는 단순하지만 오해가 많은 지표다. PUE는 compute efficiency가 아니라 facility efficiency다. 즉, PUE가 낮다고 GPU가 token을 더 잘 만든다는 뜻은 아니다. 같은 GPU cluster라도 냉각과 전력 변환 overhead가 낮으면 PUE가 낮아지고, 그만큼 facility power 중 IT load로 남는 몫이 커진다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "AI 데이터센터의 PUE는 cooling architecture에 민감하다. 고밀도 GPU rack은 공랭만으로 대응하기 어렵고, direct-to-chip liquid cooling, rear-door heat exchanger, facility water loop 같은 설계가 중요해진다. 전력 변환 효율, UPS topology, 외기 냉방 가능성, 기후, 물 사용 제약도 PUE를 좌우한다.",
                "Google은 data center별 PUE를 공개하고, Meta도 지속가능성 자료에서 PUE를 보고한다. NREL의 HPC data center 사례는 고효율 설계가 매우 낮은 PUE를 달성할 수 있음을 보여준다. 반면 Uptime Institute의 survey는 industry average PUE가 신규 hyperscale best practice보다 높을 수 있음을 보여준다. 따라서 특정 company/site가 아니라 전체 fleet을 모델링할 때는 best-in-class PUE를 그대로 쓰면 낙관 편향이 생긴다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "site-specific PUE가 공개되면 그 값을 사용한다. 공개되지 않으면 세 가지 층으로 나눈다. 첫째, hyperscaler owned new-build AI facility는 낮은 PUE band를 둘 수 있다. 둘째, colocation 또는 mixed workload facility는 중간 band가 안전하다. 셋째, 전력/냉각 retrofit 또는 급하게 ramp하는 site는 더 높은 PUE를 둘 수 있다.",
                "AI rack density가 올라간다고 항상 PUE가 나빠지는 것은 아니다. liquid cooling과 전력 설계가 잘 되면 PUE는 낮아질 수 있다. 그러나 물 사용, heat rejection, partial load operation, redundancy design 때문에 실제 연평균 PUE는 design PUE와 다를 수 있다. 따라서 PUE는 design target이 아니라 annualized operational metric으로 이해해야 한다.",
            ]),
            ("모델링 영향", [
                "PUE는 active_power_gw를 IT load로 나누는 계수다. PUE가 1.10이면 1GW facility power는 약 0.91GW IT load가 되고, PUE가 1.25이면 0.80GW IT load가 된다. PUE 차이 0.15는 대형 AI fleet에서는 수백 MW의 IT load 차이를 만들 수 있다.",
                "다만 PUE sensitivity는 active_power_gw나 inference_share보다 보통 작을 수 있다. 예를 들어 PUE 1.15와 1.20의 차이는 약 4% 수준의 IT load 차이다. 하지만 임원 보고에서는 PUE를 무시하면 facility power와 IT power를 혼동하게 되므로 반드시 명시해야 한다.",
            ]),
            ("Hallucination 위험", [
                "PUE를 GPU efficiency처럼 설명하면 안 된다. 또 company average PUE를 특정 AI cluster PUE로 확정하면 안 된다. PUE는 연평균, campus-level, site-level, design target에 따라 값이 다르므로 source의 scope를 기록해야 한다.",
            ]),
        ],
    },
    {
        "id": "A04",
        "slug": "ai_workload_share",
        "title": "ai_workload_share",
        "subtitle": "IT load 중 실제 AI training/inference cluster가 차지하는 비중",
        "sources": ["LBNL_DC", "GOOGLE_TPU_V5E", "AWS_RAINIER", "NVIDIA_GB200", "GOOGLE_IRONWOOD"],
        "starting_band": "AI-dedicated facility: 80-95%, mixed cloud facility: 50-80%, model owner attribution 불명확 시 confidence downgrade.",
        "body": [
            ("정의", [
                "ai_workload_share는 IT load 중 AI accelerator cluster가 차지하는 비중이다. IT load에는 GPU/ASIC 서버뿐 아니라 CPU 서버, storage, network, control plane, security, monitoring, data pipeline, non-AI cloud workload가 포함될 수 있다. AI 전용으로 설계된 신규 campus라면 이 비중이 높을 수 있지만, mixed cloud region이라면 낮아질 수 있다.",
                "이 변수는 active_power_gw와 inference_share 사이의 다리다. active_power_gw를 PUE로 나눠 IT load를 만든 뒤, 그중 AI workload에 해당하는 몫을 잡아야 training/inference 계산을 시작할 수 있다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "AI workload라고 해도 모두 GPU compute는 아니다. 대형 LLM 서비스에는 data ingestion, retrieval index, storage, embedding pipeline, safety classifier, logging, monitoring, load balancer, orchestration service가 필요하다. 특히 RAG와 agent workload가 커질수록 storage와 CPU-side serving infrastructure도 중요해진다. 그러나 token generation capacity를 계산할 때는 이 보조 load가 accelerator inference load와 구분되어야 한다.",
                "Google TPU v5e 문서는 training과 serving workload가 같은 TPU product 내에서도 다르게 provision될 수 있음을 보여준다. AWS Project Rainier처럼 Anthropic용 Trainium2 cluster가 명시되는 경우 AI workload share가 높다고 볼 수 있지만, 그 cluster 안에서도 training과 inference, eval, post-training이 나뉜다. NVIDIA GB200 NVL72 같은 rack-scale system은 AI accelerator load의 직접 신호지만, facility 전체의 storage/network/control plane overhead는 별도로 남는다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "AI-dedicated campus 또는 partner-specific cluster라면 ai_workload_share를 80-95%로 시작할 수 있다. 하지만 hyperscaler region 전체나 mixed-use data center라면 50-80%가 더 안전하다. 공개자료가 'AI data center'라고 말하더라도 모든 IT load가 token-generating accelerator라고 가정하면 안 된다.",
                "이 가정은 architecture maturity에 따라 바뀐다. 대형 AI cluster가 rack-scale로 통합될수록 accelerator share가 높아질 수 있지만, agentic AI와 RAG가 커질수록 storage, retrieval, networking, CPU orchestration load도 증가한다. 따라서 AI workload share를 무조건 100%에 가깝게 두는 것은 위험하다.",
            ]),
            ("모델링 영향", [
                "ai_workload_share는 inference_gw를 선형적으로 움직인다. active_power_gw와 PUE가 같을 때 ai_workload_share가 0.8에서 0.9로 오르면 AI IT load와 downstream token capacity가 12.5% 증가한다. 따라서 이 값은 생각보다 큰 민감도를 갖는다.",
                "이 가정은 memory marketing에도 중요하다. AI workload share가 높을수록 HBM과 accelerator attach가 커지고, mixed workload share가 높을수록 DDR5, enterprise SSD, network/storage infrastructure 기회도 커진다.",
            ]),
            ("Hallucination 위험", [
                "가장 큰 위험은 AI data center라는 표현을 보고 IT load 전체를 GPU inference로 간주하는 것이다. 두 번째 위험은 storage와 networking overhead를 무시하는 것이다. 세 번째 위험은 hyperscaler region 전체와 model-owner-dedicated cluster를 같은 share로 처리하는 것이다.",
            ]),
        ],
    },
    {
        "id": "A05",
        "slug": "inference_power_share",
        "title": "inference_power_share",
        "subtitle": "AI IT load 중 상용 토큰 생성에 배정되는 전력 비중",
        "sources": ["GOOGLE_IRONWOOD", "GOOGLE_TPU_V5E", "PAGED_ATTENTION", "SPLITWISE", "DISTSERVE", "AWS_RAINIER", "OPENAI_GPT41"],
        "starting_band": "2026 Base 50-60%, 2030 Base 65-80%. 2026 60%+는 fact가 아니라 company-specific bull 또는 scenario로 표현.",
        "body": [
            ("정의", [
                "inference_power_share는 AI IT load 중 inference serving에 배정되는 비중이다. inference는 학습된 모델을 이용해 prompt를 처리하고 output token을 생성하는 서비스 단계다. ChatGPT, Claude, Gemini, Copilot, Meta AI, Qwen API, Hunyuan 서비스의 traffic이 이 share를 만든다.",
                "이 값은 추론 토큰 생성량의 직접 입력이다. 그러나 공개자료에서 기업별 inference_power_share가 직접 공개되는 경우는 매우 드물다. 따라서 이 변수는 대부분 estimate 또는 scenario이며, 공식 fact처럼 표현하면 안 된다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "상용 LLM 기업이 성숙할수록 inference_power_share는 상승 압력을 받는다. 매일 들어오는 API 요청, consumer chatbot, enterprise assistant, coding agent, search answer, multimodal generation은 연중 반복되는 serving load다. 반면 training은 대규모 실험과 모델 세대 전환에 집중된다. 상용 traffic이 커지면 marginal capacity는 inference 쪽으로 많이 배정될 수 있다.",
                "하지만 inference share가 100%로 가는 것은 아니다. frontier model training, post-training, eval, safety testing, synthetic data generation, distillation, data pipeline이 계속 compute를 요구한다. 특히 frontier 경쟁을 하는 기업은 상용 inference가 커져도 training reserve를 유지한다.",
                "Google Ironwood처럼 inference-oriented accelerator 발표는 inference 수요의 방향성을 보여준다. Google TPU v5e 문서도 training과 serving provisioning 차이를 설명한다. PagedAttention, Splitwise, DistServe는 inference serving이 prefill/decode, KV cache, latency SLO, batching 문제와 연결된다는 것을 보여준다. 이런 문헌은 inference share를 직접 주지는 않지만, inference가 독립된 capacity planning 대상임을 뒷받침한다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "2026년 Base에서 50-60%를 학습용 band로 두는 이유는 상용 inference가 이미 중요하지만 frontier training과 post-training도 여전히 크기 때문이다. 2030년 Base에서 65-80%로 올리는 이유는 LLM이 제품과 workflow에 더 깊게 들어가고, inference-oriented hardware와 serving optimization이 확대될 가능성이 높기 때문이다.",
                "기업별로 조정해야 한다. OpenAI, Anthropic, Google처럼 상용 surface가 큰 기업은 inference share가 높을 수 있다. xAI처럼 빠르게 product traffic을 키우는 challenger는 초기 training-heavy에서 빠르게 inference share가 올라갈 수 있다. DeepSeek, Alibaba, Tencent는 public API와 internal app traffic의 투명성이 낮아 confidence를 낮게 두되, 중국 내 대규모 app integration 가능성을 sensitivity로 봐야 한다.",
            ]),
            ("모델링 영향", [
                "inference_power_share는 token forecast에 선형적으로 작용한다. AI IT load가 1GW이고 inference share가 60%라면 600MW가 token generation에 들어간다. share가 70%로 오르면 token capacity는 같은 조건에서 16.7% 증가한다.",
                "이 값은 memory marketing에서도 중요하다. inference share가 올라가면 HBM만이 아니라 DDR5/MRDIMM, SSD, CXL, network, power delivery 수요도 달라진다. 긴 context와 RAG가 많아질수록 KV cache, SSD retrieval, CPU-side memory attach의 중요성이 커진다.",
            ]),
            ("Hallucination 위험", [
                "2026년에 전체 AI GW의 60% 이상이 inference라고 단정하는 것은 공개자료만으로는 위험하다. inference-oriented chip 발표는 방향성이지 share fact가 아니다. 또한 inference_power_share와 utilization을 혼동하면 안 된다. inference로 배정된 capacity가 항상 100% token generation으로 전환되는 것은 아니다.",
            ]),
        ],
    },
    {
        "id": "A06",
        "slug": "training_power_share",
        "title": "training_power_share",
        "subtitle": "frontier training, post-training, eval capacity를 어떻게 남길 것인가",
        "sources": ["KAPLAN", "CHINCHILLA", "EPOCH_TRAIN", "EPOCH_POWER", "AWS_RAINIER", "GOOGLE_TPU_V5E"],
        "starting_band": "training_power_share = 1 - inference_power_share로 시작하되, frontier lab은 2026-2030에도 의미 있는 reserve 유지.",
        "body": [
            ("정의", [
                "training_power_share는 AI IT load 중 모델 weight 업데이트, pretraining, post-training, fine-tuning, RL, eval, synthetic data generation 등 학습 관련 workload에 배정되는 비중이다. 간단한 모델에서는 1 - inference_power_share로 계산하지만, 실제로는 eval/reserve/post-training을 별도 category로 두는 것이 더 정확하다.",
                "이 변수는 token generation forecast에서는 inference의 반대편처럼 보이지만, AI 기업의 전략에서는 핵심이다. training capacity는 다음 모델 세대, 성능 우위, 제품 차별화, 장기 경쟁력을 만든다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "Kaplan scaling law와 Chinchilla는 training compute allocation을 이해하는 기초 이론이다. 모델 성능은 parameter, data, compute scaling과 연결되고, compute-optimal training은 단순히 모델을 키우는 것이 아니라 학습 토큰과 모델 크기를 균형 있게 조정해야 함을 보여준다. 이 이론은 왜 frontier lab이 상용 inference가 커져도 training을 계속하는지 설명한다.",
                "Epoch AI의 frontier training compute 분석은 training run의 compute와 power demand가 계속 커질 수 있음을 보여준다. 다만 이 숫자는 연중 평균 training share가 아니라 특정 frontier run의 envelope로 읽어야 한다. peak training power와 annual average power allocation을 섞으면 잘못된 결론이 나온다.",
                "Training workload는 cluster topology에 민감하다. 수천에서 수만 accelerator가 하나의 job에 묶이면 interconnect, checkpoint, failure recovery, parallelism strategy가 중요해진다. 이런 job은 inference처럼 작은 regional fleet으로 쉽게 나누기 어렵다. 그래서 training capacity는 물리적으로도 운영적으로도 별도 planning이 필요하다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "초기 연구 중심 기업은 training share가 높다. 상용화가 진행된 기업은 inference share가 올라가지만 frontier model 경쟁을 한다면 training share를 낮게만 둘 수 없다. 2026년에는 많은 기업이 training과 inference를 모두 크게 늘리는 단계로 보는 것이 합리적이다.",
                "2030년에는 inference share 상승이 base scenario지만, training share가 사라지는 것은 아니다. 특히 reasoning model, multimodal model, agent model, long-context model은 계속 post-training과 eval compute를 요구한다. 따라서 training_power_share는 단순 잔여값이더라도 해석상 '미래 역량 투자 capacity'로 설명해야 한다.",
            ]),
            ("모델링 영향", [
                "training_power_share가 높아지면 단기 token generation capacity는 낮아진다. 그러나 장기적으로 더 효율적인 모델, 더 작은 active parameter, better routing, distillation을 가능하게 해 future tokens/sec/MW를 올릴 수 있다. 따라서 training share를 단순 비용으로만 보면 안 된다.",
                "메모리 관점에서는 training share가 높을수록 HBM bandwidth/capacity, high-end GPU allocation, scale-up/scale-out network, checkpoint storage 수요가 커진다. inference share가 높을수록 latency, KV cache, SSD retrieval, regional serving fleet 수요가 더 중요해진다.",
            ]),
            ("Hallucination 위험", [
                "training_power_share를 inference의 단순 잔여값으로 두고 설명을 생략하면 안 된다. 또 frontier training run의 peak power 전망을 연중 평균 share로 쓰면 안 된다. training에는 pretraining뿐 아니라 post-training, eval, synthetic data generation이 포함될 수 있음을 명시해야 한다.",
            ]),
        ],
    },
    {
        "id": "A07",
        "slug": "active_parameters",
        "title": "active_parameters",
        "subtitle": "Dense와 MoE에서 token당 실제 계산량을 결정하는 핵심 변수",
        "sources": ["DEEPSEEK_V3", "DEEPSEEK_R1", "QWEN3", "OPENAI_GPT41", "KAPLAN", "CHINCHILLA"],
        "starting_band": "Dense: active ~= total. MoE: official active parameter 사용. Closed model: precise number 금지, band/proxy만 사용.",
        "body": [
            ("정의", [
                "active_parameters는 token 하나를 생성할 때 실제로 활성화되어 계산에 참여하는 parameter 규모다. Dense Transformer에서는 대부분의 parameter가 매 token 계산에 관여하므로 active parameter가 total parameter와 거의 같다. 반면 MoE는 전체 expert parameter 중 일부만 token마다 activate되므로 total parameter와 active parameter가 크게 다르다.",
                "LLM inference cost를 계산할 때 active_parameters는 total_parameters보다 더 직접적이다. token당 대략적인 forward FLOPs는 2 * active_parameters라는 sanity check로 볼 수 있다. 이 식은 매우 단순하지만, closed model 숫자나 tokens/sec/MW가 물리적으로 말이 되는지 점검하는 데 유용하다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "MoE는 거대한 지식 capacity와 낮은 per-token compute cost를 동시에 노리는 구조다. DeepSeek-V3는 공식 repo에서 671B total parameters와 37B activated parameters per token을 제시한다. Qwen3 계열도 235B total / 22B active 같은 MoE 구조를 공개한다. 이 숫자들은 total parameter만 보고 inference 비용을 과대평가하면 안 된다는 강한 anchor다.",
                "Closed model은 다르다. OpenAI GPT-4.1 문서는 context window, output token limit, pricing, endpoint 등 product 정보를 제공하지만 parameter count는 공개하지 않는다. 이런 모델에 대해 인터넷 추정치를 precise fact로 쓰면 보고서 신뢰도가 무너진다. closed model은 architecture band, performance class, pricing, latency, context, benchmark proxy로 다뤄야 한다.",
                "active parameter는 KV cache와도 연결된다. token당 compute는 active parameter가 크게 좌우하지만, long context serving에서는 KV cache memory와 attention cost도 중요하다. 따라서 active parameter만으로 tokens/sec/MW를 완전히 설명할 수 없다. 모델 architecture, context length, batch shape, prefill/decode ratio를 함께 봐야 한다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "공개 MoE 모델은 official repo 또는 model card의 active parameter를 사용한다. DeepSeek, Qwen처럼 공식 자료가 있으면 confidence를 높일 수 있다. 공개 dense 모델은 total parameter를 active parameter proxy로 둔다. closed model은 단일 숫자 대신 low/base/high band를 두고, benchmark sanity layer에서만 proxy로 사용한다.",
                "active parameter band는 model family별로 잡아야 한다. 같은 회사라도 flagship reasoning model, mini model, coding model, embedding model, multimodal model이 다르다. 상용 token 대부분이 항상 가장 큰 모델에서 나온다고 가정하면 비용을 과대평가할 수 있다. 실제 서비스는 routing과 model cascade를 사용해 작은 모델이 많은 traffic을 처리할 수 있다.",
            ]),
            ("모델링 영향", [
                "active_parameters가 커지면 token당 FLOPs와 energy cost가 증가하고 tokens/sec/MW는 낮아지는 방향이다. MoE optimization이 좋아지거나 smaller model routing이 커지면 effective active parameter가 낮아져 token capacity가 증가할 수 있다.",
                "이 변수는 tokens/sec/MW를 직접 산정하지 않더라도 sanity check에 필수다. forecast가 어떤 회사의 closed flagship model에 대해 지나치게 높은 tokens/sec/MW를 가정한다면, active parameter band와 accelerator FLOPs로 역산해 모순을 찾아야 한다.",
            ]),
            ("Hallucination 위험", [
                "가장 큰 위험은 closed model parameter rumor를 fact로 쓰는 것이다. 두 번째 위험은 MoE total parameter를 active parameter처럼 사용하거나, 반대로 active parameter만 보고 memory footprint를 과소평가하는 것이다. MoE는 per-token compute와 total weight storage를 분리해야 한다.",
            ]),
        ],
    },
    {
        "id": "A08",
        "slug": "tokens_per_second_per_mw",
        "title": "tokens_per_second_per_mw",
        "subtitle": "1MW inference load가 초당 몇 token을 만들 수 있는가",
        "sources": ["PAGED_ATTENTION", "SPLITWISE", "DISTSERVE", "MS_SPLITWISE", "GOOGLE_IRONWOOD", "NVIDIA_H100", "NVIDIA_H200", "NVIDIA_GB200"],
        "starting_band": "company fact가 아니라 benchmark/proxy. 모델 크기, context, hardware, serving stack에 따라 wide band 사용.",
        "body": [
            ("정의", [
                "tokens_per_second_per_mw는 inference power 1MW가 초당 생성할 수 있는 output token 수를 의미한다. 이 변수는 token forecast의 핵심 productivity coefficient다. 하지만 단일 산업 표준값이 있는 지표가 아니다. 모델 크기, active parameter, prompt length, output length, batch shape, latency SLO, accelerator generation, memory bandwidth, interconnect, serving software에 따라 크게 달라진다.",
                "이 값은 company-level production telemetry가 공개되지 않는 한 대부분 proxy다. 따라서 보고서에서는 'benchmark calibrated assumption' 또는 'scenario coefficient'로 표시해야 한다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "LLM serving은 prefill과 decode로 나뉜다. Prefill은 prompt를 처리하는 단계이고 decode는 output token을 하나씩 생성하는 단계다. Prefill은 compute-intensive, decode는 memory/KV-cache/latency-sensitive 성격이 강하다. PagedAttention은 KV cache memory management를 개선해 batching과 throughput을 높이는 mechanism을 제시했고, Splitwise와 DistServe는 prefill과 decode를 분리해 GPU pool을 다르게 쓰는 approach를 제시한다.",
                "Hardware도 중요하다. H100/H200/B200/GB200, TPU Ironwood, Trainium/Inferentia는 compute precision, HBM capacity, memory bandwidth, interconnect, software stack이 다르다. Google Ironwood는 inference-oriented TPU로 발표되었고, NVIDIA GB200 NVL72는 rack-scale NVLink domain을 제공한다. 하지만 공식 hardware peak FLOPs를 그대로 tokens/sec/MW로 바꾸면 안 된다. production serving efficiency는 kernel, batching, scheduling, SLO, utilization에 의해 크게 낮아진다.",
                "tokens/sec/MW는 결국 effective throughput 지표다. 같은 model이라도 latency SLA를 완화하고 batch를 키우면 throughput은 올라가지만 사용자 경험이 나빠질 수 있다. enterprise real-time assistant와 offline batch summarization은 완전히 다른 efficiency를 보인다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "현실적인 접근은 세 가지 sanity check를 동시에 쓰는 것이다. 첫째, benchmark layer에서 GPU count와 model active parameter로 대략적인 tokens/sec를 역산한다. 둘째, energy layer에서 joules/token과 inference MW로 daily token을 계산한다. 셋째, service layer에서 utilization과 latency SLO를 고려해 평균 throughput을 낮춘다.",
                "closed model은 특히 wide band가 필요하다. 공개 모델 benchmark가 있다고 해서 OpenAI, Anthropic, Google production model과 동일하게 볼 수 없다. 공개 MoE 모델은 active parameter가 있으므로 상대적으로 더 좁은 band를 둘 수 있지만, 실제 serving stack과 traffic mix는 여전히 unknown이다.",
            ]),
            ("모델링 영향", [
                "tokens_per_second_per_mw는 active inference MW와 곱해져 token/day를 만든다. 따라서 이 값의 20% 차이는 token forecast의 20% 차이다. 특히 2030 forecast에서는 hardware generation과 serving efficiency improvement를 어떻게 가정하느냐가 큰 차이를 만든다.",
                "이 변수는 memory marketing에도 직접 연결된다. tokens/MW가 올라가면 같은 전력으로 더 많은 token을 만들 수 있어 HBM capacity pressure가 낮아질 수도 있지만, 실제로는 demand elasticity 때문에 더 많은 inference traffic이 생겨 전체 memory demand가 다시 증가할 수 있다.",
            ]),
            ("Hallucination 위험", [
                "가장 큰 위험은 benchmark number를 company fact처럼 쓰는 것이다. 두 번째 위험은 peak throughput을 annual average throughput으로 쓰는 것이다. 세 번째 위험은 input token과 output token을 섞는 것이다. token forecast에서 말하는 token이 input, output, total processed 중 무엇인지 반드시 명시해야 한다.",
            ]),
        ],
    },
    {
        "id": "A09",
        "slug": "utilization",
        "title": "utilization",
        "subtitle": "이론 capacity 중 실제 연평균 토큰으로 전환되는 비율",
        "sources": ["PAGED_ATTENTION", "SPLITWISE", "DISTSERVE", "GOOGLE_TPU_V5E", "AWS_INF2", "AWS_TRN2"],
        "starting_band": "2026 inference utilization 45-65%, 2030 60-80%. Peak benchmark를 평균 utilization로 사용 금지.",
        "body": [
            ("정의", [
                "utilization은 이론적으로 provisioned된 inference capacity 중 실제 token generation으로 전환되는 평균 비율이다. 여기서 평균은 보통 연중 또는 일중 평균을 의미한다. utilization이 낮다고 반드시 운영이 나쁘다는 뜻은 아니다. latency SLO, failover reserve, regional redundancy, traffic burst, model routing 때문에 의도적으로 여유 capacity를 둔다.",
                "utilization은 inference_power_share와 다르다. inference_power_share는 AI IT load 중 inference에 배정된 전력이고, utilization은 그 inference capacity가 실제 traffic 처리에 사용되는 평균 비율이다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "LLM inference traffic은 시간대, 지역, 제품 surface에 따라 변동한다. consumer chatbot은 peak/off-peak 차이가 있고, enterprise workload는 업무시간과 batch job 패턴이 다르며, coding agent는 긴 session과 tool call을 동반할 수 있다. latency SLO를 맞추려면 peak traffic을 처리할 headroom이 필요하고, 장애 시 failover를 위한 reserve도 필요하다.",
                "Serving system은 utilization과 latency 사이에서 trade-off를 가진다. batch를 크게 만들면 throughput과 utilization은 좋아지지만 first-token latency가 나빠질 수 있다. PagedAttention, Splitwise, DistServe 같은 연구는 이 trade-off를 개선하려는 방향이다. 하지만 production system이 항상 paper benchmark처럼 움직인다고 볼 수는 없다.",
                "TPU/Trainium/Inferentia/GPU 같은 accelerator는 workload shape에 따라 utilization이 다르다. training은 long-running job으로 high utilization을 목표로 할 수 있지만, inference는 SLO와 burst reserve 때문에 평균 utilization이 낮아질 수 있다. 따라서 training cluster와 inference fleet utilization을 같은 값으로 두면 안 된다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "초기 상용 서비스는 traffic 예측이 불확실하고 software stack이 안정화 중이므로 utilization을 낮게 두는 것이 안전하다. 시간이 지나 model routing, batching, KV cache management, regional load balancing이 개선되면 utilization은 상승할 수 있다. 그러나 100%에 가까운 연평균 utilization은 일반적으로 비현실적이다.",
                "2026년 45-65%, 2030년 60-80%는 학습용 band다. paid API 비중이 높고 batchable workload가 많으면 상단에 가까울 수 있다. free consumer chatbot이나 strict latency real-time workload가 많으면 하단에 가까울 수 있다.",
            ]),
            ("모델링 영향", [
                "utilization은 token forecast에 선형적으로 작용한다. tokens/sec/MW가 같더라도 utilization이 0.55에서 0.70으로 오르면 annual token은 27% 증가한다. 이 값은 software optimization과 demand smoothing의 효과를 반영하는 핵심 coefficient다.",
                "utilization 상승은 capacity 부족을 일부 완화할 수 있다. 전력 ramp가 지연되어도 batching, disaggregation, model routing이 개선되면 token output 감소를 줄일 수 있다. 그래서 grid-constrained / efficiency-upside 시나리오가 필요하다.",
            ]),
            ("Hallucination 위험", [
                "peak benchmark throughput을 utilization 100%로 연중 곱하는 것이 가장 위험하다. 두 번째 위험은 utilization을 전력 사용률처럼 해석하는 것이다. inference fleet은 전력을 쓰고 있어도 SLO reserve로 일부 capacity가 비어 있을 수 있다.",
            ]),
        ],
    },
    {
        "id": "A10",
        "slug": "attribution_rule",
        "title": "attribution_rule",
        "subtitle": "host capacity와 model-owner token을 중복 없이 귀속하는 법",
        "sources": ["OPENAI_INFRA", "OPENAI_STARGATE", "AWS_RAINIER", "AWS_TRN2", "GOOGLE_IRONWOOD", "OPENAI_GPT41"],
        "starting_band": "model-owner forecast에서는 token을 모델 소유자에게 귀속. host capacity는 별도 infrastructure layer로 기록.",
        "body": [
            ("정의", [
                "attribution_rule은 특정 전력, accelerator, cloud capacity, token output을 어느 회사 row에 귀속할지 정하는 규칙이다. LLM 생태계에서는 model owner, cloud host, infrastructure investor, product distributor가 다를 수 있기 때문에 attribution이 없으면 숫자가 쉽게 중복된다.",
                "예를 들어 Oracle data center capacity가 OpenAI를 위해 사용된다면 model-owner token forecast에서는 OpenAI capacity로 귀속할 수 있다. 그러나 infrastructure revenue 또는 data center capex 모델에서는 Oracle 또는 project owner layer에 기록해야 한다. 같은 capacity를 OpenAI와 Oracle 양쪽 token capacity에 더하면 중복 계산이다.",
            ]),
            ("전문가 관점의 운영 원리", [
                "OpenAI/Microsoft 관계는 attribution의 대표적인 난제다. Microsoft Copilot token, Azure-hosted OpenAI API, OpenAI direct ChatGPT/API token, Azure-hosted third-party model token이 섞일 수 있다. model owner 기준 forecast에서는 GPT 계열 token은 OpenAI에, Microsoft-owned Phi/MAI 계열 token은 Microsoft에, Copilot product traffic이 GPT를 호출하는 경우는 token owner와 product owner를 별도 표기해야 한다.",
                "Anthropic도 마찬가지다. AWS Project Rainier와 Trainium2 capacity는 Anthropic Claude training/inference와 연결되지만, capacity host는 AWS다. Google Cloud/Vertex route도 있을 수 있다. model-owner forecast에서는 Claude token을 Anthropic에 귀속하되, host dependency와 hardware mix는 AWS/Google layer에 기록한다.",
                "Google은 Gemini model owner와 TPU/cloud host가 같은 기업 안에 있어 attribution이 비교적 단순하지만, Google Cloud에서 third-party model을 serving하는 경우는 분리해야 한다. Meta는 open-weight Llama 다운로드와 Meta AI serving을 분리해야 한다. DeepSeek, Alibaba, Tencent는 model owner와 cloud/internal app owner가 그룹 내에 있더라도 public serving과 internal integration을 구분해야 한다.",
            ]),
            ("현실적 숫자 범위 잡기", [
                "attribution은 숫자 band가 아니라 규칙이다. 먼저 분석 목적을 정한다. model-owner token forecast라면 token을 생성하게 한 model family 소유자에게 귀속한다. cloud revenue forecast라면 hosting provider에게 귀속한다. power infrastructure forecast라면 site owner 또는 contracted power owner에게 귀속한다. memory demand forecast라면 실제 accelerator procurement와 qualification owner를 추가로 봐야 한다.",
                "각 row에는 attribution_note를 반드시 둔다. 'OpenAI model tokens hosted on Oracle/Azure/Stargate', 'Anthropic Claude tokens hosted on AWS Trainium/Bedrock and Google Vertex', 'Microsoft Copilot tokens using OpenAI model separated from Microsoft-owned Phi/MAI' 같은 문구가 필요하다.",
            ]),
            ("모델링 영향", [
                "attribution_rule은 aggregate total을 크게 바꾼다. 중복 계산을 제거하면 company별 token share가 달라지고, US vs China 비교도 달라진다. 또한 memory marketing에서는 실제 구매 의사결정자와 model owner가 다를 수 있으므로 account strategy에도 영향을 준다.",
                "이 규칙은 confidence와도 연결된다. attribution이 명확한 dedicated cluster는 confidence가 높고, multi-tenant cloud serving이나 internal app routing은 confidence가 낮다. 보고서에서는 confidence를 낮춘 이유를 'capacity가 작다'가 아니라 'attribution transparency가 낮다'로 써야 한다.",
            ]),
            ("Hallucination 위험", [
                "가장 큰 위험은 host와 model owner를 동시에 더하는 것이다. 두 번째 위험은 product owner와 model owner를 혼동하는 것이다. 세 번째 위험은 open-weight model 다운로드를 해당 회사의 hosted token generation으로 착각하는 것이다. attribution rule이 없으면 forecast는 구조적으로 부풀려진다.",
            ]),
        ],
    },
]


def source_lines(ids: list[str]) -> list[str]:
    lines = []
    for sid in ids:
        title, publisher, year, url = SOURCES[sid]
        lines.append(f"- **{sid}:** {title}, {publisher}, {year}. {url}")
    return lines


def build_markdown(item: dict) -> str:
    lines = [
        f"# {item['id']} {item['title']}",
        "",
        f"**부제:** {item['subtitle']}",
        "",
        f"**생성일:** {RUN_DATE}",
        "",
        f"**학습용 starting band:** {item['starting_band']}",
        "",
        f"> {COMMON_WARNING}",
        "",
        "## Executive Summary",
        "",
        f"`{item['title']}`는 LLM token capacity simulation에서 숫자 정합성을 좌우하는 핵심 가정입니다. 이 문서는 해당 가정이 무엇을 의미하는지, 어떤 물리·운영·경제 원리에 의해 결정되는지, 공개자료를 어떻게 읽어야 하는지, 그리고 모델에 넣을 때 어떤 hallucination 위험을 피해야 하는지를 전문가 관점에서 정리합니다.",
        "",
    ]
    for heading, paragraphs in item["body"]:
        lines.append(f"## {heading}")
        lines.append("")
        for para in paragraphs:
            lines.append(para)
            lines.append("")
    lines.extend(
        [
            "## 실무 체크리스트",
            "",
            "- 이 값은 fact, estimate, proxy, scenario 중 무엇인가?",
            "- 단위가 GW, MW, TWh/year, tokens/sec, tokens/day 중 무엇인지 명확한가?",
            "- 회사 공식자료가 직접 말한 숫자인가, 우리가 계산한 숫자인가?",
            "- 값이 10% 변하면 2030 Base forecast가 얼마나 움직이는가?",
            "- low/base/high band와 confidence가 같이 기록되어 있는가?",
            "- 새로운 공식 source가 나오면 어떤 replacement path로 교체할 것인가?",
            "",
            "## 권장 모델 필드",
            "",
            f"- assumption_id: `{item['id']}_{item['slug'].upper()}`",
            f"- field_name: `{item['slug']}`",
            "- derivation_type: Fact / Estimate / Proxy / Scenario 중 하나",
            "- confidence: High / Medium / Low",
            "- replacement_path: 공식 source가 나오면 교체할 경로",
            "",
            "## 참고자료",
            "",
            *source_lines(item["sources"]),
            "",
            "## 저작권 및 사용 메모",
            "",
            "이 문서는 외부 자료의 전문을 복제하지 않고, 모델링과 학습에 필요한 범위에서 해석과 요약을 제공합니다. 보고서에 수치를 사용할 때는 원문 source를 다시 열어 title, publisher, date, URL, 숫자 정의를 확인해야 합니다.",
        ]
    )
    return "\n".join(lines)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    run.font.name = "Apple SD Gothic Neo"
    run.font.size = Pt(10)


def build_docx(item: dict, path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Apple SD Gothic Neo"
    styles["Normal"].font.size = Pt(10)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Apple SD Gothic Neo"
        styles[style_name].font.color.rgb = RGBColor(15, 23, 42)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{item['id']} {item['title']}")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Apple SD Gothic Neo"
    run.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(item["subtitle"])

    doc.add_heading("Executive Summary", level=1)
    add_para(doc, f"학습용 starting band: {item['starting_band']}")
    add_para(doc, COMMON_WARNING)
    add_para(doc, f"`{item['title']}`는 LLM token capacity simulation에서 숫자 정합성을 좌우하는 핵심 가정입니다. 이 문서는 해당 가정이 무엇을 의미하는지, 어떤 물리·운영·경제 원리에 의해 결정되는지, 공개자료를 어떻게 읽어야 하는지, 그리고 모델에 넣을 때 어떤 hallucination 위험을 피해야 하는지를 전문가 관점에서 정리합니다.")

    for heading, paragraphs in item["body"]:
        doc.add_heading(heading, level=1)
        for para in paragraphs:
            add_para(doc, para)

    doc.add_heading("실무 체크리스트", level=1)
    checklist = [
        "이 값은 fact, estimate, proxy, scenario 중 무엇인가?",
        "단위가 GW, MW, TWh/year, tokens/sec, tokens/day 중 무엇인지 명확한가?",
        "회사 공식자료가 직접 말한 숫자인가, 우리가 계산한 숫자인가?",
        "값이 10% 변하면 2030 Base forecast가 얼마나 움직이는가?",
        "low/base/high band와 confidence가 같이 기록되어 있는가?",
        "새로운 공식 source가 나오면 어떤 replacement path로 교체할 것인가?",
    ]
    for line in checklist:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("권장 모델 필드", level=1)
    fields = [
        f"assumption_id: {item['id']}_{item['slug'].upper()}",
        f"field_name: {item['slug']}",
        "derivation_type: Fact / Estimate / Proxy / Scenario 중 하나",
        "confidence: High / Medium / Low",
        "replacement_path: 공식 source가 나오면 교체할 경로",
    ]
    for line in fields:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("참고자료", level=1)
    for line in source_lines(item["sources"]):
        add_para(doc, line.replace("**", ""))

    doc.add_heading("저작권 및 사용 메모", level=1)
    add_para(doc, "이 문서는 외부 자료의 전문을 복제하지 않고, 모델링과 학습에 필요한 범위에서 해석과 요약을 제공합니다. 보고서에 수치를 사용할 때는 원문 source를 다시 열어 title, publisher, date, URL, 숫자 정의를 확인해야 합니다.")
    doc.save(path)


def build_index() -> str:
    lines = [
        "# Assumption Textbook Index",
        "",
        f"생성일: {RUN_DATE}",
        "",
        "이 폴더는 LLM token capacity simulation의 10개 핵심 가정을 각각 독립 리포트로 공부하기 위한 자료입니다.",
        "",
        "| ID | 가정 | 공부 질문 | Markdown | Word |",
        "|---|---|---|---|---|",
    ]
    for item in ASSUMPTIONS:
        md = f"{item['id']}_{item['slug']}.md"
        docx = f"{item['id']}_{item['slug']}.docx"
        lines.append(f"| {item['id']} | {item['title']} | {item['subtitle']} | {md} | outputs/reports/assumptions/{docx} |")
    lines.extend(
        [
            "",
            "## 추천 학습 순서",
            "",
            "1. A01-A02로 전력 capacity의 상한과 active 전환을 이해합니다.",
            "2. A03-A04로 facility power와 AI IT load의 차이를 이해합니다.",
            "3. A05-A06으로 training/inference allocation을 공부합니다.",
            "4. A07-A09로 모델 구조, serving efficiency, utilization을 공부합니다.",
            "5. A10으로 company attribution을 정리해 중복 계산을 제거합니다.",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    generated = []
    for item in ASSUMPTIONS:
        name = f"{item['id']}_{item['slug']}"
        md_path = DOCS / f"{name}.md"
        docx_path = OUT / f"{name}.docx"
        md_path.write_text(build_markdown(item), encoding="utf-8")
        build_docx(item, docx_path)
        generated.append((str(md_path), str(docx_path)))
    (DOCS / "README.md").write_text(build_index(), encoding="utf-8")
    print({"status": "PASS", "reports": len(generated), "outputs": generated})


if __name__ == "__main__":
    main()
