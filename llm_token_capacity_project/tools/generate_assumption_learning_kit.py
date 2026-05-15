"""Generate a Korean learning kit for LLM token capacity assumptions.

The deck and markdown are educational artifacts for building intuition before
changing model coefficients. They intentionally distinguish official facts,
derived assumptions, and calibration exercises.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT / "outputs" / "reports"
RUN_DATE = "2026-05-15"


SOURCES = [
    {
        "id": "SRC_GOOGLE_IRONWOOD",
        "title": "Ironwood: The first Google TPU for the age of inference",
        "publisher": "Google Cloud",
        "date": "2025-04-09",
        "url": "https://blog.google/innovation-and-ai/infrastructure-and-cloud/google-cloud/ironwood-tpu-age-of-inference/",
        "lesson": "TPU pod, inference-oriented accelerator, HBM, perf/W 학습 anchor",
    },
    {
        "id": "SRC_OPENAI_STARGATE",
        "title": "OpenAI, Oracle, and SoftBank expand Stargate with five new AI data center sites",
        "publisher": "OpenAI",
        "date": "2025-09-23",
        "url": "https://openai.com/index/five-new-stargate-sites/",
        "lesson": "planned GW와 active AI IT load를 분리해야 하는 이유",
    },
    {
        "id": "SRC_DEEPSEEK_V3",
        "title": "DeepSeek-V3 GitHub / Technical Report",
        "publisher": "DeepSeek",
        "date": "2024-12-26",
        "url": "https://github.com/deepseek-ai/DeepSeek-V3",
        "lesson": "MoE total parameter와 active parameter 분리",
    },
    {
        "id": "SRC_QWEN3",
        "title": "Qwen3 GitHub",
        "publisher": "Alibaba Qwen",
        "date": "2025",
        "url": "https://github.com/QwenLM/Qwen3",
        "lesson": "MoE 모델의 active parameter와 thinking/non-thinking serving 차이",
    },
    {
        "id": "SRC_NVIDIA_H100",
        "title": "NVIDIA H100 Tensor Core GPU",
        "publisher": "NVIDIA",
        "date": "accessed 2026-05-14",
        "url": "https://www.nvidia.com/en-us/data-center/h100/",
        "lesson": "GPU TDP와 rack/facility power는 다르다는 점",
    },
    {
        "id": "SRC_ARXIV_ENERGY_TOKEN",
        "title": "Energy Use of AI Inference: Efficiency Pathways and Test-Time Compute",
        "publisher": "arXiv",
        "date": "2025",
        "url": "https://arxiv.org/abs/2509.20241",
        "lesson": "query/token 단위 에너지 모델링과 test-time compute 민감도",
    },
]


MODULES = [
    {
        "id": "M1",
        "name": "전력 capacity",
        "question": "계약된 GW 중 실제로 AI IT load가 되는 비중은?",
        "concept": "Contracted/planned capacity는 최대 경계값이고, active power는 인허가, 변전, 냉각, 랙 설치, GPU 수급을 통과한 운영 capacity입니다.",
        "starter_band": "2026 active/contracted: 15-45%, 2030: 45-80%",
        "watch": "planned GW를 곧바로 inference GW로 쓰는 오류",
        "sources": "SRC_OPENAI_STARGATE",
    },
    {
        "id": "M2",
        "name": "PUE와 AI workload share",
        "question": "시설 전력에서 GPU/ASIC 서버가 실제로 먹는 전력은?",
        "concept": "PUE는 facility power를 IT load로 바꾸는 계수이고, AI workload share는 IT load 중 AI cluster가 차지하는 비중입니다.",
        "starter_band": "PUE: 1.10-1.30, AI workload share: 70-95%",
        "watch": "facility MW, IT MW, GPU board power를 같은 숫자로 취급",
        "sources": "SRC_GOOGLE_IRONWOOD; SRC_NVIDIA_H100",
    },
    {
        "id": "M3",
        "name": "Training / inference mix",
        "question": "전체 AI IT load 중 inference가 몇 %인가?",
        "concept": "상용 서비스가 커질수록 inference 비중은 상승하지만, frontier training과 post-training capacity는 계속 남습니다.",
        "starter_band": "2026 Base: 50-60%, 2030 Base: 65-80%",
        "watch": "2026년 60%+를 fact로 표현",
        "sources": "SRC_OPENAI_STARGATE; SRC_GOOGLE_IRONWOOD",
    },
    {
        "id": "M4",
        "name": "모델 파라미터와 MoE",
        "question": "total parameter와 token당 active parameter가 어떻게 다른가?",
        "concept": "Dense 모델은 대부분의 parameter가 매 token 계산에 관여하지만, MoE는 일부 expert만 activate되어 active parameter가 훨씬 낮습니다.",
        "starter_band": "Dense: active ~= total, MoE: active/total 3-20%",
        "watch": "MoE total parameter만 보고 inference 비용을 과대평가",
        "sources": "SRC_DEEPSEEK_V3; SRC_QWEN3",
    },
    {
        "id": "M5",
        "name": "GPU/ASIC mix",
        "question": "NVIDIA GPU, TPU, Trainium, 자체 ASIC이 토큰 효율에 주는 영향은?",
        "concept": "동일 MW라도 hardware generation, memory bandwidth, interconnect, serving software에 따라 tokens/sec/MW가 크게 달라집니다.",
        "starter_band": "GPU-heavy: 낮은 custom efficiency, TPU/ASIC-heavy: 높은 serving efficiency 가능",
        "watch": "GPU 수만 세고 accelerator 세대와 utilization을 무시",
        "sources": "SRC_GOOGLE_IRONWOOD; SRC_NVIDIA_H100",
    },
    {
        "id": "M6",
        "name": "Tokens/sec/MW",
        "question": "1MW inference load가 초당 몇 token을 만들 수 있는가?",
        "concept": "tokens/sec/MW는 모델 크기, active parameter, batch size, context length, KV cache, quantization, speculative decoding의 합성 결과입니다.",
        "starter_band": "closed model: proxy band만 사용, open MoE: active parameter로 sanity check",
        "watch": "benchmark 값을 회사별 실제 성능처럼 확정",
        "sources": "SRC_ARXIV_ENERGY_TOKEN; SRC_DEEPSEEK_V3",
    },
    {
        "id": "M7",
        "name": "Utilization",
        "question": "이론 capacity 중 실제 토큰으로 변환되는 비율은?",
        "concept": "실제 utilization은 traffic shape, latency SLA, batch 가능성, regional placement, failover reserve 때문에 100%가 될 수 없습니다.",
        "starter_band": "2026: 45-65%, 2030: 60-80%",
        "watch": "peak throughput을 연중 평균 throughput으로 사용",
        "sources": "SRC_ARXIV_ENERGY_TOKEN",
    },
    {
        "id": "M8",
        "name": "Attribution",
        "question": "누구의 전력 capacity를 누구의 token으로 귀속할 것인가?",
        "concept": "AWS, Oracle, Google Cloud capacity는 hosting이고, token owner는 Claude, GPT, Gemini 등 model owner 기준으로 귀속합니다.",
        "starter_band": "capacity owner != model owner인 경우 attribution rule 필수",
        "watch": "OpenAI/Microsoft, Anthropic/AWS, OpenAI/Oracle 중복 계산",
        "sources": "SRC_OPENAI_STARGATE",
    },
]


def add_textbox(slide, x, y, w, h, text, size=16, bold=False, color=RGBColor(23, 37, 84), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    p = frame.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Apple SD Gothic Neo"
    return box


def add_title(slide, title, subtitle=""):
    add_textbox(slide, 0.55, 0.35, 11.9, 0.45, title, 23, True, RGBColor(15, 23, 42))
    if subtitle:
        add_textbox(slide, 0.58, 0.83, 11.4, 0.3, subtitle, 10.5, False, RGBColor(71, 85, 105))


def add_footer(slide, note="Source: project source register; numbers are learning bands unless explicitly marked as official facts."):
    add_textbox(slide, 0.55, 7.08, 12.0, 0.25, note, 7.5, False, RGBColor(100, 116, 139))


def add_bar(slide, x, y, w, h, fill):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    return shape


def add_card(slide, x, y, w, h, title, body, fill=RGBColor(248, 250, 252), accent=RGBColor(37, 99, 235)):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(203, 213, 225)
    add_bar(slide, x, y, 0.08, h, accent)
    add_textbox(slide, x + 0.18, y + 0.12, w - 0.32, 0.28, title, 12, True, RGBColor(15, 23, 42))
    add_textbox(slide, x + 0.18, y + 0.47, w - 0.32, h - 0.52, body, 9.5, False, RGBColor(51, 65, 85))


def set_bg(slide, color=RGBColor(255, 255, 255)):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_bullets(slide, x, y, w, h, bullets, size=11):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    for idx, item in enumerate(bullets):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Apple SD Gothic Neo"
        p.font.color.rgb = RGBColor(51, 65, 85)


def write_markdown(path: Path) -> None:
    lines = [
        "# Assumption Learning Kit",
        "",
        f"생성일: {RUN_DATE}",
        "",
        "이 문서는 LLM token capacity simulation의 가정을 현실적인 숫자로 만들기 위한 학습 도구입니다. 숫자를 바로 외우기보다, 어떤 질문을 던지고 어떤 단위로 검증해야 하는지 익히는 것이 목적입니다.",
        "",
        "## 먼저 외울 핵심 문장",
        "",
        "- 계약 GW는 token capacity가 아니라 상한선입니다.",
        "- active GW는 전력망, 변전, 냉각, 랙, accelerator 수급을 통과한 운영 capacity입니다.",
        "- inference GW는 active AI IT load 중 inference에 배정된 부분입니다.",
        "- MoE 모델은 total parameter보다 active parameter가 token cost에 더 직접적입니다.",
        "- tokens/sec/MW는 hardware, model, software, traffic shape가 합쳐진 결과입니다.",
        "- benchmark는 sanity check이지 회사별 실제 성능 fact가 아닙니다.",
        "",
        "## 1페이지 계산 지도",
        "",
        "```text",
        "contracted GW",
        "  -> active GW",
        "  -> IT load GW = active GW / PUE",
        "  -> AI IT load GW = IT load GW * AI workload share",
        "  -> inference GW = AI IT load GW * inference share",
        "  -> inference MW",
        "  -> tokens/day = MW * tokens/sec/MW * utilization * 86,400",
        "```",
        "",
        "## Assumption Modules",
        "",
    ]
    for module in MODULES:
        lines.extend(
            [
                f"### {module['id']}. {module['name']}",
                "",
                f"**핵심 질문:** {module['question']}",
                "",
                f"**개념:** {module['concept']}",
                "",
                f"**초기 학습 band:** {module['starter_band']}",
                "",
                f"**주의할 오류:** {module['watch']}",
                "",
                f"**관련 source:** {module['sources']}",
                "",
                "**연습:**",
                "",
                "1. 이 가정이 10% 올라가면 token forecast가 얼마나 바뀌는지 계산한다.",
                "2. 이 가정을 fact로 바꿀 수 있는 공식 source가 있는지 찾는다.",
                "3. 공식 source가 없다면 estimate/proxy/scenario 중 어디에 속하는지 표시한다.",
                "",
            ]
        )
    lines.extend(
        [
            "## 숫자를 현실적으로 만드는 5단계",
            "",
            "1. 먼저 단위를 고정한다: GW, MW, tokens/sec/MW, tokens/day, annual tokens.",
            "2. 그 다음 physical boundary를 확인한다: active <= contracted, inference+training=100%.",
            "3. 모델 구조를 확인한다: dense인지 MoE인지, active parameter가 공개되어 있는지.",
            "4. benchmark와 비교한다: GPU count 방식과 tokens/MW 방식이 크게 벌어지는지.",
            "5. 마지막으로 문구를 조정한다: fact, estimate, proxy, scenario를 분리한다.",
            "",
            "## Calibration Exercises",
            "",
            "### Exercise A: Planned GW와 Active GW",
            "",
            "OpenAI/Stargate 같은 planned capacity 발표를 보고 active_power_gw를 바로 같게 두지 말고, 2026/2030 deployment ratio를 별도로 둡니다.",
            "",
            "질문:",
            "",
            "- 발표 문구가 planned, contracted, committed, operational 중 무엇인가?",
            "- 데이터센터가 energized 되었는가?",
            "- GPU rack delivery와 cooling이 같이 확인되었는가?",
            "- training/inference workload가 시작됐다는 문구가 있는가?",
            "",
            "### Exercise B: MoE Parameter",
            "",
            "DeepSeek-V3처럼 671B total / 37B active가 공개된 모델은 total parameter만 보고 token cost를 계산하면 안 됩니다.",
            "",
            "질문:",
            "",
            "- active parameter가 token당 기준인가?",
            "- MTP, speculative decoding, quantization이 serving에 영향을 주는가?",
            "- context length가 길어질 때 prefill/KV cache 비용이 늘어나는가?",
            "",
            "### Exercise C: Benchmark 괴리",
            "",
            "main forecast와 benchmark reference가 50% 이상 차이 나면 숫자가 틀렸다고 단정하지 말고 먼저 원인을 분해합니다.",
            "",
            "- 모델 active parameter proxy가 너무 낮거나 높은가?",
            "- GPU count 추정이 facility power와 accelerator board power를 혼동했는가?",
            "- utilization을 peak 기준으로 둔 것은 아닌가?",
            "- TPU/ASIC custom efficiency를 NVIDIA GPU 기준으로 비교하고 있지는 않은가?",
            "",
            "## Source Reading List",
            "",
            "| source_id | title | publisher | date | learning use |",
            "|---|---|---|---|---|",
        ]
    )
    for src in SOURCES:
        lines.append(f"| {src['id']} | [{src['title']}]({src['url']}) | {src['publisher']} | {src['date']} | {src['lesson']} |")
    lines.extend(
        [
            "",
            "## Next Study Loop",
            "",
            "매주 한 가정만 골라서 다음 순서로 공부합니다.",
            "",
            "1. 공식 source 2개 읽기",
            "2. 숫자 3개만 추출하기",
            "3. fact/estimate/proxy/scenario로 분류하기",
            "4. 현재 모델의 값과 비교하기",
            "5. 차이가 나면 assumption_change_log.md에 후보 변경으로 기록하기",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def build_ppt(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    navy = RGBColor(15, 23, 42)
    blue = RGBColor(37, 99, 235)
    green = RGBColor(16, 185, 129)
    amber = RGBColor(245, 158, 11)
    red = RGBColor(220, 38, 38)
    pale_blue = RGBColor(239, 246, 255)
    pale_green = RGBColor(236, 253, 245)
    pale_amber = RGBColor(255, 251, 235)
    pale_gray = RGBColor(248, 250, 252)

    slide = prs.slides.add_slide(blank)
    set_bg(slide, RGBColor(248, 250, 252))
    add_textbox(slide, 0.75, 0.75, 10.9, 0.5, "LLM Token Capacity", 20, True, blue)
    add_textbox(slide, 0.75, 1.35, 11.3, 0.9, "가정을 현실적인 숫자로 만드는 학습 키트", 33, True, navy)
    add_textbox(slide, 0.78, 2.55, 9.6, 0.55, "전력, GPU/ASIC, 모델 구조, serving 효율, 수요 attribution을 한 장씩 분해해서 공부합니다.", 16, False, RGBColor(71, 85, 105))
    add_card(slide, 0.8, 4.25, 3.6, 1.15, "목표", "숫자를 외우는 것이 아니라, 숫자가 현실적인지 물어보는 순서를 익힌다.", pale_blue, blue)
    add_card(slide, 4.75, 4.25, 3.6, 1.15, "산출물", "MD 워크북 + PPT 학습자료 + 기존 simulation과 연결되는 assumption register.", pale_green, green)
    add_card(slide, 8.7, 4.25, 3.6, 1.15, "원칙", "Fact, Estimate, Proxy, Scenario를 절대 섞지 않는다.", pale_amber, amber)
    add_footer(slide, f"Generated {RUN_DATE}")

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "한 장으로 보는 계산 지도", "토큰 forecast는 전력에서 바로 나오지 않고 여러 assumption gate를 통과한다.")
    steps = [
        ("1", "Contracted GW", "계약/계획 capacity"),
        ("2", "Active GW", "실제로 energize된 운영 capacity"),
        ("3", "AI IT load", "PUE와 workload share 적용"),
        ("4", "Inference GW", "training/inference split 적용"),
        ("5", "Tokens/day", "tokens/sec/MW와 utilization 적용"),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = 0.65 + i * 2.55
        add_bar(slide, x, 2.0, 1.0, 1.0, [blue, green, amber, RGBColor(14, 165, 233), RGBColor(124, 58, 237)][i])
        add_textbox(slide, x, 2.17, 1.0, 0.35, num, 20, True, RGBColor(255, 255, 255), PP_ALIGN.CENTER)
        add_textbox(slide, x - 0.25, 3.25, 1.5, 0.3, title, 12, True, navy, PP_ALIGN.CENTER)
        add_textbox(slide, x - 0.45, 3.65, 1.9, 0.5, body, 9, False, RGBColor(71, 85, 105), PP_ALIGN.CENTER)
        if i < 4:
            add_textbox(slide, x + 1.15, 2.27, 0.8, 0.2, "→", 24, True, RGBColor(148, 163, 184), PP_ALIGN.CENTER)
    add_card(slide, 1.0, 5.15, 11.3, 0.8, "핵심 감각", "각 단계는 곱셈 구조이므로 작은 가정 차이가 2030 token forecast에서는 큰 차이로 증폭된다.", pale_gray, blue)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Assumption map", "공부할 가정은 8개 모듈로 나누면 관리가 쉬워진다.")
    for i, module in enumerate(MODULES):
        x = 0.7 + (i % 4) * 3.15
        y = 1.35 + (i // 4) * 2.35
        fill = [pale_blue, pale_green, pale_amber, pale_gray][i % 4]
        accent = [blue, green, amber, RGBColor(100, 116, 139)][i % 4]
        add_card(slide, x, y, 2.75, 1.7, f"{module['id']} {module['name']}", module["question"], fill, accent)
    add_footer(slide)

    concept_slides = [
        ("전력 capacity: planned와 active를 분리", MODULES[0], ["발표된 GW는 대부분 상한선 또는 목표치", "active GW는 site energization, transformer, cooling, rack, accelerator delivery의 결과", "2030 forecast는 deployment ratio가 핵심 민감도"]),
        ("PUE와 AI workload share", MODULES[1], ["Facility power를 IT load로 바꾸려면 PUE 필요", "IT load 전체가 GPU cluster는 아님", "GPU board TDP와 rack/facility power는 다른 단위"]),
        ("Training / inference mix", MODULES[2], ["Commercial LLM traffic이 커질수록 inference 비중 상승", "frontier training, post-training, eval capacity는 계속 필요", "2026년 60%+는 fact가 아니라 scenario로 취급"]),
        ("모델 구조와 MoE", MODULES[3], ["Dense는 active ~= total에 가까움", "MoE는 token당 일부 expert만 activate", "DeepSeek-V3 같은 공개 MoE는 active parameter가 중요한 anchor"]),
        ("GPU/ASIC mix", MODULES[4], ["NVIDIA GPU, TPU, Trainium, custom ASIC은 tokens/MW가 다름", "HBM 용량/대역폭과 interconnect가 serving efficiency에 영향", "GPU count만으로 token capacity를 확정할 수 없음"]),
        ("Tokens/sec/MW", MODULES[5], ["모델 크기, context, batching, quantization, speculative decoding의 합성 결과", "benchmark는 company fact가 아니라 proxy", "closed model은 band와 confidence로 다뤄야 함"]),
        ("Utilization", MODULES[6], ["Peak throughput과 연평균 throughput은 다름", "latency SLA와 failover reserve가 utilization을 낮춤", "traffic shape가 좋을수록 batching 효율 상승"]),
        ("Attribution", MODULES[7], ["host capacity와 model owner token을 분리", "OpenAI/Microsoft, Anthropic/AWS, OpenAI/Oracle 중복 주의", "계약 당사자와 모델 토큰 owner가 다를 수 있음"]),
    ]
    for idx, (title, module, bullets) in enumerate(concept_slides):
        slide = prs.slides.add_slide(blank)
        set_bg(slide)
        add_title(slide, title, module["question"])
        add_card(slide, 0.8, 1.35, 4.0, 1.4, "개념", module["concept"], [pale_blue, pale_green, pale_amber, pale_gray][idx % 4], [blue, green, amber, RGBColor(100, 116, 139)][idx % 4])
        add_card(slide, 5.05, 1.35, 3.15, 1.4, "초기 학습 band", module["starter_band"], pale_gray, green)
        add_card(slide, 8.45, 1.35, 3.95, 1.4, "주의할 오류", module["watch"], pale_amber, red)
        add_textbox(slide, 0.85, 3.45, 2.6, 0.3, "공부 포인트", 15, True, navy)
        add_bullets(slide, 0.9, 3.92, 5.6, 1.9, bullets, 12)
        add_textbox(slide, 7.05, 3.45, 2.8, 0.3, "직접 해볼 질문", 15, True, navy)
        add_bullets(
            slide,
            7.1,
            3.92,
            5.0,
            1.9,
            [
                "이 값은 fact인가, estimate인가?",
                "10% 바꾸면 forecast가 얼마나 움직이는가?",
                "공식 source로 교체할 수 있는가?",
            ],
            12,
        )
        add_footer(slide, f"Related source ids: {module['sources']}")

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "시나리오를 현실적으로 만드는 법", "Bull/Base/Bear는 느낌이 아니라 세 가지 축의 조합이어야 한다.")
    axes = [
        ("전력 operational deploy", "인허가, 변전, 냉각, 랙 설치, GPU delivery 속도", blue),
        ("MoE / serving 최적화", "active parameter, quantization, batching, speculative decoding", green),
        ("Inference 비중 상승", "상용 traffic, API/Copilot/agent workload 확대", amber),
    ]
    for i, (name, body, color) in enumerate(axes):
        add_card(slide, 0.9 + i * 4.05, 1.45, 3.55, 1.35, name, body, [pale_blue, pale_green, pale_amber][i], color)
    add_card(slide, 0.9, 3.65, 3.55, 1.25, "Bear", "전력 투입 지연 + 효율 개선 둔화 + inference mix 상승 지연", pale_gray, red)
    add_card(slide, 4.95, 3.65, 3.55, 1.25, "Base", "공식 발표를 staged deployment로 반영하고 점진적 효율 개선", pale_gray, blue)
    add_card(slide, 9.0, 3.65, 3.55, 1.25, "Bull", "빠른 energization + 강한 serving 효율 + 상용 inference 급증", pale_gray, green)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Calibration exercise", "숫자가 이상해 보이면 틀렸다고 말하기 전에 어디서 벌어졌는지 분해한다.")
    add_textbox(slide, 0.85, 1.35, 4.0, 0.32, "예시 질문", 15, True, navy)
    add_bullets(
        slide,
        0.9,
        1.85,
        5.4,
        3.6,
        [
            "planned GW를 active GW로 바로 썼는가?",
            "PUE 적용 전후를 혼동했는가?",
            "MoE total parameter만 보고 token cost를 계산했는가?",
            "peak throughput을 연평균으로 사용했는가?",
            "benchmark proxy를 company fact로 표현했는가?",
        ],
        12,
    )
    add_card(slide, 6.9, 1.45, 4.95, 1.1, "1단계", "단위: GW → MW → tokens/sec → tokens/day → annual tokens", pale_blue, blue)
    add_card(slide, 6.9, 2.85, 4.95, 1.1, "2단계", "경계: active <= contracted, training + inference = 100%", pale_green, green)
    add_card(slide, 6.9, 4.25, 4.95, 1.1, "3단계", "모델: dense/MoE, active params, context, serving software 확인", pale_amber, amber)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "학습 루프", "매주 한 가정만 깊게 보면 모델 전체의 신뢰도가 올라간다.")
    loop = [
        ("1", "Source 읽기", "공식 source 2개"),
        ("2", "숫자 추출", "숫자 3개만"),
        ("3", "분류", "fact/estimate/proxy/scenario"),
        ("4", "모델 비교", "현재 값과 차이 확인"),
        ("5", "로그 기록", "assumption/source log 업데이트"),
    ]
    for i, (num, title, body) in enumerate(loop):
        x = 0.8 + i * 2.5
        add_bar(slide, x, 2.0, 0.78, 0.78, [blue, green, amber, RGBColor(14, 165, 233), RGBColor(124, 58, 237)][i])
        add_textbox(slide, x, 2.13, 0.78, 0.25, num, 16, True, RGBColor(255, 255, 255), PP_ALIGN.CENTER)
        add_textbox(slide, x - 0.3, 3.05, 1.4, 0.26, title, 12, True, navy, PP_ALIGN.CENTER)
        add_textbox(slide, x - 0.45, 3.45, 1.7, 0.42, body, 9.5, False, RGBColor(71, 85, 105), PP_ALIGN.CENTER)
    add_card(slide, 1.0, 5.2, 11.3, 0.8, "다음 추천 주제", "첫 주는 active_power_gw와 operational deployment ratio부터 공부하세요. forecast 민감도가 가장 크고, hallucination 위험도 가장 높습니다.", pale_blue, blue)
    add_footer(slide)

    slide = prs.slides.add_slide(blank)
    set_bg(slide)
    add_title(slide, "Source reading list", "공부는 source_id 단위로 짧게, 반복적으로 한다.")
    for i, src in enumerate(SOURCES):
        x = 0.75 + (i % 2) * 6.1
        y = 1.25 + (i // 2) * 1.55
        add_card(slide, x, y, 5.65, 1.08, src["id"], f"{src['publisher']} · {src['date']}\n{src['lesson']}", [pale_blue, pale_green, pale_amber][i % 3], [blue, green, amber][i % 3])
    add_footer(slide, "Full URLs are in docs/assumption_learning_kit.md")

    prs.save(path)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    write_markdown(DOCS / "assumption_learning_kit.md")
    build_ppt(OUT / "assumption_learning_kit_kr.pptx")
    print(
        {
            "status": "PASS",
            "markdown": str(DOCS / "assumption_learning_kit.md"),
            "pptx": str(OUT / "assumption_learning_kit_kr.pptx"),
        }
    )


if __name__ == "__main__":
    main()
